"""Class, Confidence and Rubric Band: read from the start of the model's notes into columns."""

import pytest
from openpyxl import load_workbook

from asics_agent.agents.sections import base
from asics_agent.agents.sections.base import AnswerReply, parse_class_line
from asics_agent.llm import CallFailed

REAL_CALL_JSON = base.call_json


@pytest.mark.parametrize(
    ("notes", "expected"),
    [
        (
            "CLASS: A1; CONFIDENCE: High; BAND: Full marks. ALSO: BWSSB-02",
            ("A1", "High", "Full marks"),
        ),
        ("class:a2 ;  confidence : medium; band : 5 - partial", ("A2", "Medium", "5 - partial")),
        ("  CLASS = needs review, CONFIDENCE = LOW, BAND = n/a", ("Needs review", "Low", "n/a")),
        (
            "CLASS: M; CONFIDENCE: Low; BAND: (a) yes; (b) ?; (c) no CONFIRM: ask the ULG",
            ("M", "Low", "(a) yes; (b) ?; (c) no"),
        ),
        ("CLASS: R; CONFIDENCE: Low; BAND: none\nRTI-TARGET: the PIO", ("R", "Low", "none")),
    ],
)
def test_the_three_values_are_read_from_the_start_of_the_notes(notes, expected):
    *values, problems = parse_class_line(notes)
    assert tuple(values) == expected and problems == []


def test_a_value_that_is_not_allowed_is_left_blank_not_guessed():
    answer_class, confidence, band, problems = parse_class_line(
        "CLASS: A3; CONFIDENCE: Very high; BAND: Full"
    )
    assert (answer_class, confidence, band) == ("", "", "Full")
    assert any("CLASS 'A3' is not one of A1, A2, R, P, M, Needs review" in p for p in problems)
    assert any("CONFIDENCE 'Very high' is not one of High, Medium, Low" in p for p in problems)


def test_missing_values_are_reported():
    assert parse_class_line("CLASS: A1; CONFIDENCE: High")[2:] == (
        "",
        ["BAND is missing from the notes, so that column was left blank."],
    )
    *values, problems = parse_class_line("The plan says so. CLASS: A1; CONFIDENCE: High; BAND: x")
    assert values == ["", "", ""]
    assert len(problems) == 1 and "don't start with CLASS" in problems[0]
    assert parse_class_line("")[:3] == ("", "", "")


def test_a2_can_never_be_high():
    *_, problems = parse_class_line("CLASS: A2; CONFIDENCE: High; BAND: x")
    assert problems == ["Class A2 is never High confidence. Check the Confidence."]
    assert parse_class_line("CLASS: A1; CONFIDENCE: High; BAND: x")[3] == []
    assert parse_class_line("CLASS: A2; CONFIDENCE: Medium; BAND: x")[3] == []


@pytest.fixture
def answers_with(monkeypatch, run_steps):
    """Step 2 answers where the model's notes start with `prefix`."""

    def run(prefix: str):
        def call_json(*args, **kwargs):
            reply, urls, notes = REAL_CALL_JSON(*args, **kwargs)
            if isinstance(reply, AnswerReply):
                reply = reply.model_copy(update={"notes": prefix + reply.notes})
            return reply, urls, notes

        run_steps("sources")
        monkeypatch.setattr(base, "call_json", call_json)
        return run_steps("answers")

    return run


def test_the_columns_are_filled_and_the_notes_are_kept(answers_with):
    result = answers_with("CLASS: A1; CONFIDENCE: High; BAND: Full marks ALSO: SWB-01. ")
    answer = next(a for a in result["answers"] if a.question_id == "UPD 1b" and a.citation_id)
    assert (answer.answer_class, answer.confidence, answer.rubric_band) == (
        "A1",
        "High",
        "Full marks",
    )
    assert (
        "CLASS: A1; CONFIDENCE: High; BAND: Full marks ALSO: SWB-01. The plan documents ULG participation."
        in answer.notes
    )
    assert "Reviewer:" not in answer.notes

    ws = load_workbook(result["output_workbook"])["Scoring - Sample City A"]
    header = [c.value for c in ws[1]]
    first = header.index("Max Score") + 1
    assert header[first : first + 3] == ["Class", "Confidence", "Rubric Band"]
    rows = {
        (r[header.index("Parastatal")], r[header.index("Question ID")]): r
        for r in ws.iter_rows(min_row=2, values_only=True)
    }
    row = next(r for r in rows.values() if r[header.index("Citation ID")] == answer.citation_id)
    assert row[first : first + 3] == ("A1", "High", "Full marks")
    options = {v.formula1 for v in ws.data_validations.dataValidation}
    assert '"A1,A2,R,P,M,Needs review"' in options and '"High,Medium,Low"' in options


def test_a_reply_without_the_prefix_leaves_the_columns_blank_and_says_why(answers_with):
    result = answers_with("")
    answer = next(a for a in result["answers"] if a.question_id == "UPD 1b" and a.citation_id)
    assert (answer.answer_class, answer.confidence, answer.rubric_band) == ("", "", "")
    assert "Reviewer: The notes don't start with CLASS, CONFIDENCE and BAND" in answer.notes
    assert "The plan documents ULG participation." in answer.notes


def test_a2_with_high_confidence_is_flagged_in_the_notes(answers_with):
    result = answers_with("CLASS: A2; CONFIDENCE: High; BAND: Part A. ")
    answer = next(a for a in result["answers"] if a.question_id == "UPD 1b" and a.citation_id)
    assert (answer.answer_class, answer.confidence) == ("A2", "High")  # kept as the model gave it
    assert "Reviewer: Class A2 is never High confidence. Check the Confidence." in answer.notes


def test_a_failed_call_has_no_reviewer_noise(monkeypatch, run_steps):
    def failing(*args, **kwargs):
        raise CallFailed("the research took too long")

    run_steps("sources")
    monkeypatch.setattr(base, "call_json", failing)
    answers = run_steps("answers")["answers"]
    assert answers and all("Reviewer:" not in a.notes and not a.answer_class for a in answers)
