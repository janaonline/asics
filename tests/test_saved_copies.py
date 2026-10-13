"""Pages the research team saved by hand for links that block automated tools."""

import hashlib
import os
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asics_agent.agents import common
from asics_agent.agents.final_qa import qa_sources
from asics_agent.models import AccessCheck, Source
from asics_agent.practice import FILLER, INDIA_CODE, PLAN, PLAN_TEXT
from asics_agent.sources_workbook import CITATIONS_SHEET, read_sources_workbook

TEAM_NOTE = "Verified by the research team from a saved copy"
PLAN_PAGE = f"<html><body><p>{PLAN_TEXT}</p></body></html>"


def block(workbook: str, url: str, use: str = "Yes") -> str:
    """Make the row for `url` look as if the link checker had been refused, as for a
    government website that blocks automated tools. Returns its Citation ID."""
    wb = load_workbook(workbook)
    ws = wb[CITATIONS_SHEET]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if url in str(row[header.index("Official URL")].value) and (
            row[header.index("Parastatal")].value == "SWB"
        ):
            row[header.index("Verification Status")].value = "Human Verification Required"
            row[header.index("Accessibility")].value = "Not human-verified"
            row[header.index("Authority Score")].value = 0
            row[header.index("Use for Answers?")].value = use
            citation_id = row[header.index("Citation ID")].value
            break
    wb.save(workbook)
    evidence = Path(workbook).parent / "evidence"
    for saved in evidence.glob(f"{hashlib.sha1(url.encode()).hexdigest()[:16]}.*"):
        saved.unlink()  # no automatic copy of this page
    return citation_id


def save_copy(workbook: str, name: str, content: str | bytes = PLAN_PAGE) -> Path:
    path = Path(workbook).parent / "evidence" / "manual" / name
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(content if isinstance(content, bytes) else content.encode())
    os.utime(path, (time.time() - 100, time.time() - 100))  # older than any later automatic copy
    return path


def read(workbook: str):
    _, citations, issues = read_sources_workbook(Path(workbook), Path(workbook).parent / "evidence")
    return {c.url: c for c in citations if c.parastatal_id == "SWB"}, [i.message for i in issues]


@pytest.fixture
def refused(monkeypatch):
    """The link checker is refused by PLAN's website, as by a blocking government site."""
    real = common.check_url

    def check(url, client, cache_dir, user_agent):
        if url == PLAN:
            return AccessCheck(
                url=url,
                verification_status="Human Verification Required",
                reasons=["HTTP 403: the website refused the automated check."],
                checked_at="2026-10-12T10:00:00",
            )
        return real(url, client, cache_dir, user_agent)

    monkeypatch.setattr(common, "check_url", check)


@pytest.mark.parametrize(
    "name", ["{id}.html", "{id}.htm", "{id}.txt", "{id}.pdf", "{id_lower}.html"]
)
def test_a_saved_copy_verifies_a_blocked_link(run_steps, monkeypatch, name):
    workbook = run_steps("sources")["output_workbook"]
    citation_id = block(workbook, PLAN)
    if name.endswith(".pdf"):
        monkeypatch.setattr("asics_agent.links.manual.pdf_to_text", lambda data: PLAN_TEXT)
        content = b"%PDF-1.4 (a PDF with a text layer)"
    else:
        content = PLAN_TEXT if name.endswith(".txt") else PLAN_PAGE
    copy = save_copy(workbook, name.format(id=citation_id, id_lower=citation_id.lower()), content)

    sources, problems = read(workbook)
    row = sources[PLAN]
    assert row.verification_status == "Verified" and row.accessibility == "Public"
    assert row.team_saved_copy and row.snapshot_path == str(copy)
    assert "recorded its responses" in Path(row.content_path).read_text(encoding="utf-8")
    assert TEAM_NOTE in row.notes and f"evidence/manual/{copy.name}" in row.notes
    assert not problems


def test_step2_answers_from_a_team_saved_copy(run_steps):
    workbook = run_steps("sources")["output_workbook"]
    save_copy(workbook, f"{block(workbook, PLAN)}.html")

    answers = {(a.parastatal_id, a.question_id): a for a in run_steps("answers")["answers"]}
    answer = answers[("SWB", "UPD 1b")]
    assert answer.status == "Answered — Verified Source" and answer.citation_url == PLAN


@pytest.mark.parametrize(
    ("url", "text", "message"),
    [
        (PLAN, "Only a few words were saved.", "almost no readable text"),
        (INDIA_CODE, f"Title: Water Act, 1964. Type: Act. {FILLER}", "India Code"),
        (PLAN, f"Access denied. {FILLER}", "blocked page"),
    ],
)
def test_a_saved_copy_that_fails_the_content_checks_is_not_verified(run_steps, url, text, message):
    workbook = run_steps("sources")["output_workbook"]
    citation_id = block(workbook, url)
    save_copy(workbook, f"{citation_id}.txt", text)

    row = read(workbook)[0][url]
    assert row.verification_status == "Human Verification Required" and not row.team_saved_copy
    assert row.content_path is None
    problems = read(workbook)[1]
    assert any(citation_id in p and message in p for p in problems)


def test_use_for_answers_stays_as_the_team_set_it(run_steps):
    workbook = run_steps("sources")["output_workbook"]
    save_copy(workbook, f"{block(workbook, PLAN, use='No')}.html")

    row = read(workbook)[0][PLAN]
    assert row.verification_status == "Verified" and row.use_for_answers is False


def test_a_file_that_matches_no_citation_id_is_reported(run_steps):
    workbook = run_steps("sources")["output_workbook"]
    save_copy(workbook, "BWSSB-99.pdf")
    save_copy(workbook, "notes.docx")

    problems = read(workbook)[1]
    assert any("BWSSB-99.pdf matches no Citation ID" in p for p in problems)
    assert any("notes.docx was ignored" in p for p in problems)


def test_a_failed_recheck_does_not_undo_a_team_saved_copy(run_steps, refused):
    first = run_steps("sources")
    citation_id = block(first["output_workbook"], PLAN)
    copy = save_copy(first["output_workbook"], f"{citation_id}.html")

    again = run_steps("sources")
    row = next(s for s in again["final_sources"] if s.url == PLAN and s.parastatal_id == "SWB")
    assert row.verification_status == "Verified" and row.team_saved_copy
    assert row.snapshot_path == str(copy) and Path(row.content_path).exists()
    assert row.authority_score == 10  # assessed like any other verified source
    assert "Automated check: HTTP 403" in row.notes and TEAM_NOTE in row.notes
    assert not any(
        "saved copy" in i.message for i in again["final_issues"] if i.severity == "error"
    )

    # The workbook keeps it, and its Saved Copy link opens the team's file.
    ws = load_workbook(again["output_workbook"])[CITATIONS_SHEET]
    header = [c.value for c in ws[1]]
    saved = next(
        r[header.index("Saved Copy")] for r in ws.iter_rows(min_row=2) if r[0].value == citation_id
    )
    assert saved.hyperlink and saved.hyperlink.target.endswith(f"manual/{copy.name}")
    assert read(again["output_workbook"])[0][PLAN].team_saved_copy


def test_a_recheck_that_works_uses_the_fresh_automatic_copy(run_steps):
    first = run_steps("sources")
    citation_id = block(first["output_workbook"], PLAN)
    copy = save_copy(first["output_workbook"], f"{citation_id}.html")

    again = run_steps("sources")  # PLAN's website opens normally this time
    row = next(s for s in again["final_sources"] if s.url == PLAN and s.parastatal_id == "SWB")
    assert row.verification_status == "Verified" and not row.team_saved_copy
    assert row.snapshot_path != str(copy) and "fresh copy" in row.notes
    assert "Automated check: Opened the exact URL" in row.notes
    assert not read(again["output_workbook"])[0][PLAN].team_saved_copy  # the fresh copy stays


def test_final_qa_applies_the_content_rules_to_saved_copies(tmp_path):
    text = tmp_path / "BWSSB-07.txt"
    text.write_text("SOURCE URL: x\n\nA few words only.", encoding="utf-8")
    source = Source(
        citation_id="BWSSB-07",
        parastatal_id="SWB",
        title="Act",
        url=INDIA_CODE,
        verification_status="Verified",
        accessibility="Public",
        team_saved_copy=True,
        content_path=str(text),
        checked_at="2026-10-12T10:00:00",
        notes="Verified by the research team from a saved copy.",
    )
    messages = [i.message for i in qa_sources([source], set()) if i.severity == "error"]
    assert any("saved copy" in m for m in messages)
