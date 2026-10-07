import time

from openpyxl import load_workbook

from asics_agent.runner import BatchRun

CITIES = {"sample_city_a": "Sample City A", "sample_city_b": "Sample City B"}
INPUTS = {"only_questions": ["UPD 1"]}


def wait_for(batch, *statuses, timeout=60):
    deadline = time.time() + timeout
    while batch.status not in statuses and time.time() < deadline:
        time.sleep(0.05)
    return batch.status


def rows(path):
    ws = load_workbook(path)["All Cities"]
    return {r[0]: r for r in ws.iter_rows(min_row=5, values_only=True)}


def test_step1_batch_reviews_all_cities_together(services):
    batch = BatchRun(services, CITIES, INPUTS, label="Sources", step="sources").start()
    assert wait_for(batch, "waiting_for_review", "failed") == "waiting_for_review"
    assert set(batch.reviews) == set(CITIES)

    batch.submit_review(
        {"sample_city_a": {"action": "approve"}, "sample_city_b": {"action": "stop"}}
    )
    assert wait_for(batch, "finished", "failed") == "finished", batch.error
    index = rows(batch.summary_path)
    assert index["Sample City A"][1] == "Done"
    assert index["Sample City B"][1].startswith("Stopped at the review step")
    # One Sources workbook per city, in each city's own folder.
    for name in CITIES.values():
        assert list((services.settings.outputs_dir / "cities" / name).glob("*_Sources.xlsx"))


def test_step2_batch_writes_one_answers_workbook_per_city(services):
    first = BatchRun(services, CITIES, {**INPUTS, "auto_approve": True}, step="sources").start()
    assert wait_for(first, "finished", "failed") == "finished", first.error

    batch = BatchRun(services, {**CITIES, "atlantis": "Atlantis"}, INPUTS, step="answers").start()
    assert wait_for(batch, "finished", "failed") == "finished", batch.error
    assert batch.runs["atlantis"].blocked  # not in the register: reported, doesn't block others
    index = rows(batch.summary_path)
    assert index["Sample City A"][1] == index["Sample City B"][1] == "Done"
    assert index["Atlantis"][1].startswith("Not run")
    workbooks = sorted(
        p.name
        for p in (services.settings.outputs_dir / "cities").glob(
            "*/answers/*/ASICS_Parastatal_*.xlsx"
        )
    )
    assert workbooks == [
        "ASICS_Parastatal_Sample_City_A_Phase2.xlsx",
        "ASICS_Parastatal_Sample_City_B_Phase2.xlsx",
    ]
