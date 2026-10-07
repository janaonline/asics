from openpyxl import load_workbook

from asics_agent.cities import PARASTATALS_SHEET, load_register, write_register
from asics_agent.config import get_settings
from asics_agent.practice import practice_register


def test_real_register_loads_without_errors():
    configs, issues = load_register(get_settings().city_register)
    assert configs, "the city register should have at least one city ready to run"
    assert not [i for i in issues if i.severity == "error"], [i.message for i in issues]


def test_mistakes_are_reported_in_plain_language_per_city(tmp_path):
    path = practice_register(tmp_path / "register.xlsx")
    wb = load_workbook(path)
    ws = wb[PARASTATALS_SHEET]
    ws.append(
        [
            "Sample City A",
            "SWB",
            "Duplicate",
            "Water supply board",
            "Yes",
            "Yes",
            "Yes",
            "Yes",
            "Yes",
            "",
        ]
    )  # same ID twice in one city
    ws.append(["Sample City B", "BUS", "Bus company", "Bus service", "", "", "", "", "", ""])
    ws.append(["Sample Citty B", "X", "Typo city", "Other", "", "", "", "", "", ""])
    wb.save(path)

    configs, issues = load_register(path)
    messages = {i.message: i for i in issues}
    assert all(i.audience == "team" for i in issues)
    dup = next(m for m in messages if "already has a parastatal with the ID SWB" in m)
    assert messages[dup].ref == "Sample City A" and messages[dup].severity == "error"
    assert any("Type of BUS must be one of" in m for m in messages)
    assert any('"Sample Citty B" isn\'t on the Cities sheet' in m for m in messages)
    # Both cities still load with their valid rows.
    assert [p.id for p in configs["sample_city_a"].parastatals] == ["SWB"]
    assert [p.id for p in configs["sample_city_b"].parastatals] == ["SWB"]


def test_blank_yes_no_answers_use_defaults_and_say_so(tmp_path):
    path = practice_register(tmp_path / "register.xlsx")
    wb = load_workbook(path)
    ws = wb[PARASTATALS_SHEET]
    for column in range(5, 9):
        ws.cell(row=2, column=column).value = None
    wb.save(path)
    configs, issues = load_register(path)
    water = configs["sample_city_a"].parastatals[0]
    assert (water.revenue_raising, water.has_annual_budget) == (False, True)
    assert sum("left blank" in i.message for i in issues) == 4


def test_register_round_trips(tmp_path):
    configs, _ = load_register(practice_register(tmp_path / "a.xlsx"))
    again, _ = load_register(write_register(tmp_path / "b.xlsx", list(configs.values())))
    assert again == configs
