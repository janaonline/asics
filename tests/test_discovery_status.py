"""Parastatal discovery has three outcomes: eligible (ticked), needs_review (listed, unticked),
and clear fails (notes only)."""

import json
import time

from langchain_core.messages import AIMessage
from langchain_core.outputs import ChatGeneration, ChatResult
from openpyxl import load_workbook

from asics_agent.agents.parastatal_discovery import Discovered, DiscoverReply
from asics_agent.agents.sources_master import build_sources_graph
from asics_agent.practice import FakeClaude, _plain
from asics_agent.reporting import review_flags
from asics_agent.runner import PipelineRun
from asics_agent.sources_workbook import PARASTATALS_SHEET

INPUTS = {"city": "sample_city_a", "only_questions": ["UPD 1"]}
NO_ONE = "no parastatals will be researched"


def item(pid, name, kind, **extra):
    return {
        "id": pid,
        "name": name,
        "type": kind,
        "why_included": extra.pop("why_included", "Delivers the service."),
        **extra,
    }


def discovering(services, parastatals, notes=""):
    """Make the discovery agent return exactly this reply (no web evidence)."""

    class Scripted(FakeClaude):
        def _generate(self, messages, stop=None, run_manager=None, **kwargs):
            if "TASK: discover_parastatals" in _plain(messages[0].content):
                reply = {"parastatals": parastatals, "notes": notes}
                message = AIMessage(
                    content=json.dumps(reply), response_metadata={"stop_reason": "end_turn"}
                )
                return ChatResult(generations=[ChatGeneration(message=message)])
            return super()._generate(messages, stop, run_manager, **kwargs)

    services.llm = Scripted()
    return services


def run_step1(services, **extra):
    return build_sources_graph(services).invoke({**INPUTS, "auto_approve": True, **extra})


def team_messages(result):
    return [i.message for i in result["issues"] if i.audience == "team"]


def test_reply_model_defaults_to_eligible_and_never_ticks_an_unknown_status():
    base = {"id": "A", "name": "A"}
    assert Discovered(**base).status == "eligible"  # field missing: backward compatible
    assert Discovered(**base, status=None).status == "eligible"
    assert Discovered(**base, status="needs_review").status == "needs_review"
    assert Discovered(**base, status=" Needs-Review ").status == "needs_review"
    assert Discovered(**base, status="maybe").status == "needs_review"  # not auto-ticked
    assert DiscoverReply.model_validate({"parastatals": [base]}).parastatals[0].status == "eligible"


def test_status_decides_include(services):
    discovering(
        services,
        [
            item("STC", "Sample Transport", "transport_corporation", status="eligible"),
            item("SDA", "Sample Development Authority", "development_authority"),  # no status
            item(
                "SMR",
                "Sample Metro",
                "transport_corporation",
                status="needs_review",
                why_included="NEEDS REVIEW (E1): builds but may not operate the metro.",
            ),
        ],
    )
    found = {p.id: p for p in run_step1(services)["parastatals"]}
    assert found["STC"].include and found["SDA"].include
    assert not found["SMR"].include
    assert found["SMR"].found_by == "Agent"
    assert found["SMR"].why_included.startswith("NEEDS REVIEW (E1)")
    assert found["SMR"].notes  # the evidence note is kept
    assert found["SWB"].include  # the register's parastatal is untouched


def test_needs_review_without_the_label_gets_one(services):
    discovering(services, [item("SMR", "Sample Metro", "other", status="needs_review")])
    found = {p.id: p for p in run_step1(services)["parastatals"]}
    assert found["SMR"].why_included == "NEEDS REVIEW: Delivers the service."


def test_unproven_domain_is_kept_for_review_not_dropped(services):
    reason = "NEEDS REVIEW (domain proof pending): sudaonline.org says it is constituted under s.4"
    discovering(
        services,
        [
            item(
                "SUDA",
                "Sample Urban Development Authority",
                "development_authority",
                status="needs_review",
                why_included=reason,
            )
        ],
    )
    found = {p.id: p for p in run_step1(services)["parastatals"]}
    assert not found["SUDA"].include
    assert found["SUDA"].why_included == reason


def test_candidates_for_all_three_categories_are_kept_as_needs_review(services):
    discovering(
        services,
        [
            item(pid, name, kind, status="needs_review")
            for pid, name, kind in [
                ("DEV", "Development Authority", "development_authority"),
                ("WAT", "Water Board Two", "water_supply_board"),
                ("BUS", "Bus Corporation", "transport_corporation"),
            ]
        ],
    )
    found = {p.id: p for p in run_step1(services)["parastatals"]}
    assert {"DEV", "WAT", "BUS"} <= set(found)
    assert not any(found[i].include for i in ("DEV", "WAT", "BUS"))


def test_team_entry_wins_over_a_needs_review_duplicate(services):
    discovering(
        services,
        [
            item(
                "SWB",
                "Water Board",
                "water_supply_board",
                status="needs_review",
                why_included="NEEDS REVIEW (E2): may be owned by the city government.",
            )
        ],
    )
    swb = [p for p in run_step1(services)["parastatals"] if p.id == "SWB"]
    assert len(swb) == 1
    assert swb[0].include and swb[0].found_by == "Your team"


def test_warns_when_nothing_will_be_researched(services):
    review = item("SMR", "Sample Metro", "transport_corporation", status="needs_review")
    discovering(services, [review])
    # the register's SWB is always ticked, so limit the run to the one needing review
    result = run_step1(services, only_parastatals=["SMR"])
    messages = team_messages(result)
    assert any(NO_ONE in m and "NEEDS REVIEW" in m and "city register" in m for m in messages)
    assert any(m.startswith("Sample City A: no parastatals") for m in messages)

    discovering(services, [item("STC", "Sample Transport", "transport_corporation")])
    assert not any(NO_ONE in m for m in team_messages(run_step1(services)))  # SWB and STC


def test_review_screen_lists_needs_review_unticked_and_ticking_researches_it(services):
    discovering(
        services,
        [
            item("STC", "Sample Transport", "transport_corporation"),
            item(
                "SMR",
                "Sample Metro",
                "transport_corporation",
                status="needs_review",
                why_included="NEEDS REVIEW (E5): constituted but not shown to function.",
            ),
        ],
    )
    run = PipelineRun(services, INPUTS, "test", step="sources").start()
    deadline = time.time() + 30
    while run.status not in ("waiting_for_review", "failed") and time.time() < deadline:
        time.sleep(0.05)
    assert run.status == "waiting_for_review"
    shown = {p["id"]: p for p in run.review["parastatals"]}
    assert set(shown) == {"SWB", "STC", "SMR"}
    assert shown["SMR"]["include"] is False
    assert shown["SMR"]["why_included"].startswith("NEEDS REVIEW (E5)")
    assert shown["STC"]["include"] is True
    assert "SMR" in run.review["applicability"]

    run.submit_review({"action": "approve"})  # approving leaves NEEDS REVIEW out
    deadline = time.time() + 30
    while run.status not in ("finished", "failed") and time.time() < deadline:
        time.sleep(0.05)
    assert run.status == "finished", run.error
    assert {p.id: p.include for p in run.result["parastatals"]}["SMR"] is False


def test_ticking_a_needs_review_agency_includes_it(services):
    discovering(
        services,
        [item("SMR", "Sample Metro", "transport_corporation", status="needs_review")],
    )
    run = PipelineRun(services, INPUTS, "test", step="sources").start()
    deadline = time.time() + 30
    while run.status not in ("waiting_for_review", "failed") and time.time() < deadline:
        time.sleep(0.05)
    assert run.status == "waiting_for_review"
    run.submit_review({"action": "edit", "parastatals": ["SWB", "SMR"]})
    deadline = time.time() + 30
    while run.status not in ("finished", "failed") and time.time() < deadline:
        time.sleep(0.05)
    assert run.status == "finished", run.error
    assert {p.id: p.include for p in run.result["parastatals"]}["SMR"] is True
    assert any(s.parastatal_id == "SMR" for s in run.result["final_sources"])


def test_workbook_writes_include_no_for_needs_review(services):
    discovering(
        services,
        [
            item(
                "SMR",
                "Sample Metro",
                "transport_corporation",
                status="needs_review",
                why_included="NEEDS REVIEW (E6): jurisdiction unclear.",
            )
        ],
    )
    result = run_step1(services)
    ws = load_workbook(result["output_workbook"])[PARASTATALS_SHEET]
    header = [c.value for c in ws[1]]
    rows = {r[header.index("Parastatal ID")].value: r for r in ws.iter_rows(min_row=2)}
    assert rows["SMR"][header.index("Include?")].value == "No"
    assert rows["SMR"][header.index("Why Included")].value.startswith("NEEDS REVIEW (E6)")
    assert rows["SWB"][header.index("Include?")].value == "Yes"
    assert not any(s.parastatal_id == "SMR" for s in result["final_sources"])  # not researched


def test_pre_operational_eligible_agency_is_kept_ticked(services):
    why = "Category 3 (builds and will operate the metro). E1-E6 pass. PRE-OPERATIONAL: under construction"
    discovering(
        services,
        [
            item(
                "GMRC",
                "Metro Rail Company",
                "transport_corporation",
                status="eligible",
                why_included=why,
            )
        ],
    )
    found = {p.id: p for p in run_step1(services)["parastatals"]}
    assert found["GMRC"].include
    assert "PRE-OPERATIONAL: under construction" in found["GMRC"].why_included


def test_review_flags_show_what_the_team_should_see_before_ticking():
    needs = {"include": False, "why_included": "NEEDS REVIEW (E2): owned by the city government."}
    assert review_flags(needs) == ["NEEDS REVIEW"]
    both = {
        "include": True,
        "why_included": "E1-E6 pass. PRE-OPERATIONAL: trial runs. DOMAIN: sudaonline.org - proof pending",
    }
    assert review_flags(both) == ["PRE-OPERATIONAL", "DOMAIN"]
    domain = {
        "include": False,
        "why_included": "NEEDS REVIEW (domain proof pending): site says s.4",
    }
    assert review_flags(domain) == ["NEEDS REVIEW", "DOMAIN"]
    assert review_flags({"include": True, "why_included": "Delivers water."}) == []
    # unticked by the team, no NEEDS REVIEW prefix: not labelled
    assert review_flags({"include": False, "why_included": "Delivers water."}) == []
    assert review_flags({}) == []
