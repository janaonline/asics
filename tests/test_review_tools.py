from pathlib import Path

import httpx
from openpyxl import load_workbook

from asics_agent.links.text import locate_quote
from asics_agent.practice import PAGES, PLAN
from asics_agent.recheck import recheck_workbook
from asics_agent.reporting import team_issues


def test_locate_quote_finds_pdf_page_and_heading():
    text = "[page 1]\nIntro\n[page 2]\n## Section 16. Plans\nThe Board shall\nconsult the ULG."
    assert locate_quote(text, "the board shall consult the ULG") == (2, "Section 16. Plans")
    assert locate_quote(text, "not in the text") == (None, None)


def test_team_only_sees_messages_about_questions_in_the_run(run_steps):
    run_steps("sources")
    result = run_steps("answers")
    refs = {i.ref for i in team_issues(result["final_issues"])}
    assert "UPD 1a" in refs  # scale problem in a question that was run
    assert "DPG d" not in refs  # problem in a question that wasn't


def test_recheck_reports_changed_pages(run_steps):
    run_steps("sources")
    workbook = run_steps("answers")["output_workbook"]

    def handler(request):
        url = str(request.url)
        if url.endswith("/robots.txt"):
            return httpx.Response(404)
        if url == PLAN:  # the plan page has since been rewritten
            return httpx.Response(200, html="<p>" + "different words " * 200 + "</p>")
        return httpx.Response(200, html=PAGES[url]) if url in PAGES else httpx.Response(404)

    client = httpx.Client(transport=httpx.MockTransport(handler))
    out, summary = recheck_workbook(Path(workbook), client, "test")
    ws = load_workbook(out)["Scoring - Sample City A"]
    header = [c.value for c in ws[1]]
    rows = {r[header.index("Question ID")]: r for r in ws.iter_rows(min_row=2, values_only=True)}
    quote_col = header.index("Quote Still on Page?")
    assert rows["UPD 1a"][quote_col] == "Yes"
    assert rows["UPD 1b"][quote_col].startswith("No")
    assert summary["quotes no longer on the page"] == 2  # UPD 1b for both parastatals
