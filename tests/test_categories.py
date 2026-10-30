"""Categories: each is a settings file; all run the same Steps 1, 2 and 3."""

import dataclasses

import pytest

from asics_agent.categories import (
    ULG_UNIT_ID,
    agent_for,
    category_settings,
    load_bank,
    load_categories,
)
from asics_agent.config import _base_settings, get_settings


def test_parastatal_settings_are_unchanged():
    s = get_settings()
    base = _base_settings()
    assert s.category == "parastatal" and s.unit == "parastatal"
    assert s.question_bank == base.question_bank
    assert s.outputs_dir == base.outputs_dir / "Parastatal"  # outputs/Parastatal/cities/…
    assert s.shared_rules == "shared-rules" and s.category_title == "Parastatal"
    for role, agent in [
        ("discover", "parastatal-discovery"),
        ("profile", "website-profiler"),
        ("sources", "citation-builder"),
        ("answer", "answer-writer"),
        ("score", "question-scorer"),
    ]:
        assert agent_for(s, role) == agent


def test_every_category_is_complete():
    found, issues = load_categories(get_settings().agent_setup_dir)
    assert set(found) == {"parastatal", "ulg"} and issues == []


def test_ulg_has_its_own_bank_agents_and_folders():
    u = category_settings(get_settings(), "ulg")
    assert u.unit == "ulg" and u.scoring_code == "UPD"  # scoring files keep the UPD code
    assert u.outputs_dir.name in {"ULG", "UPD"}  # UPD: work saved before the rename
    assert u.category_title == "ULG"
    assert agent_for(u, "answer") == "city-answer-writer" and u.shared_rules == "shared-rules-city"
    questions, issues = load_bank(u)
    assert len(questions) > 90 and questions["UPD 3"].children == ["UPD 3a", "UPD 3b", "UPD 3c"]
    assert not [i for i in issues if i.severity == "error"]


@pytest.fixture
def ulg_services(tmp_path):
    from asics_agent.practice import practice_services

    ulg = dataclasses.replace(
        category_settings(get_settings(), "ulg"),
        outputs_dir=tmp_path,
        outputs_root=tmp_path,
        scoring_dir=tmp_path / "scoring",
    )
    return practice_services(ulg)


def test_ulg_runs_steps_1_2_3(ulg_services, tmp_path):
    from asics_agent.agents.answers_master import build_answers_graph
    from asics_agent.agents.sources_master import build_sources_graph
    from asics_agent.cities import load_register
    from asics_agent.scoring.ai import run_ai_scoring
    from asics_agent.scoring.phase2 import phase2_status
    from asics_agent.scoring.skeleton import build_skeleton, gather_units
    from asics_agent.scoring.template import read_template

    s = ulg_services.settings
    inputs = {"city": "sample_city_a", "only_questions": ["UPD 3"], "auto_approve": True}
    before = len(ulg_services.llm.calls)
    sources = build_sources_graph(ulg_services).invoke(inputs)
    assert [p.id for p in sources["parastatals"]] == [ULG_UNIT_ID]  # the ULG
    assert "discover_parastatals" not in ulg_services.llm.calls[before:]  # nothing to discover
    build_answers_graph(ulg_services).invoke(inputs)
    status = phase2_status(s, "Sample City A")
    assert status.ready and status.answers.name.startswith("ASICS_ULG_")
    check_ulg_headers(status)

    register, _ = load_register(s.city_register)
    questions, issues = load_bank(s)
    units, _ = gather_units(s)
    path = s.scoring_dir / "templates" / "upd.xlsx"
    build_skeleton(questions, units, path, "UPD", issues)  # the scoring code stays UPD
    report = run_ai_scoring(
        ulg_services, read_template(path), "UPD", register, codes=["UPD3a"], workers=1
    )
    assert report.cities == ["Sample City A"]
    assert report.scored + report.needs_person == 1  # one row: the ULG


def check_ulg_headers(status):
    """The ULG workbooks say "ULG" where the Parastatal ones say "Parastatal"; old files load."""
    from openpyxl import load_workbook

    from asics_agent.sources_workbook import read_sources_workbook
    from asics_agent.workbook import is_scoring_sheet

    sources = load_workbook(status.sources)
    assert "ULG" in sources.sheetnames and "Parastatals" not in sources.sheetnames
    assert [c.value for c in sources["ULG"][1]][:2] == ["ULG ID", "ULG Name"]
    assert sources["Citation Sheet"]["B1"].value == "ULG"
    assert sources["Coverage"]["A1"].value == "ULG"
    answers = load_workbook(status.answers)
    sheet = next(ws for ws in answers.worksheets if is_scoring_sheet(ws.title))
    assert sheet["A1"].value == "ULG" and "ULG Type" in [c.value for c in sheet[1]]
    assert "Parastatal" not in {c.value for ws in sources.worksheets for c in ws[1]}

    # The same file with the old headers (what is on disk today) reads the same.
    evidence = status.sources.parent / "evidence"
    new = read_sources_workbook(status.sources, evidence)[:2]
    sources["ULG"].title = "Parastatals"
    for ws in sources.worksheets:
        for cell in ws[1]:
            if cell.value in ("ULG ID", "ULG Name"):
                cell.value = cell.value.replace("ULG", "Parastatal")
            elif cell.value == "ULG":
                cell.value = "Parastatal"
    old = status.sources.with_name("old_headers.xlsx")
    sources.save(old)
    assert read_sources_workbook(old, evidence)[:2] == new


def test_parastatal_workbooks_keep_their_headers(run_steps):
    from openpyxl import load_workbook

    out = run_steps("sources")["output_workbook"]
    wb = load_workbook(out)
    assert "Parastatals" in wb.sheetnames and "ULG" not in wb.sheetnames
    assert [c.value for c in wb["Parastatals"][1]][:2] == ["Parastatal ID", "Parastatal Name"]
    assert wb["Citation Sheet"]["B1"].value == "Parastatal"
