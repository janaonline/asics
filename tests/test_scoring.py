"""Scoring: draft workbook, interns' copies, merge, reading scores, results workbook."""

import pytest
from openpyxl import load_workbook

from asics_agent.models import CityConfig, Parastatal, Question
from asics_agent.scoring import human
from asics_agent.scoring.recalc import recalculate, soffice
from asics_agent.scoring.scores import build_results, read_scores
from asics_agent.scoring.skeleton import Unit, build_skeleton, check_questions
from asics_agent.scoring.template import read_template


def _q(qid, tag, score, applicability="Common.", methodology="", parent=None, children=()):
    return Question(
        id=qid,
        raw_id=qid,
        row=1,
        pillar="DPG",
        pillar_name="DPG",
        text=qid,
        tag=tag,
        max_score=score,
        applicability=applicability,
        methodology=methodology,
        parent_id=parent,
        children=list(children),
    )


QUESTIONS = {
    "DPG 1": _q("DPG 1", "MQ", 10, children=["DPG 1a", "DPG 1b"]),
    "DPG 1a": _q("DPG 1a", "SQ", 5, parent="DPG 1"),
    "DPG 1b": _q("DPG 1b", "SQ", 5, "Water supply board", parent="DPG 1"),
    "DPG 5a": _q(
        "DPG 5a",
        "SQ",
        1,
        methodology="Part A: Availability — 0.5\nScore 0.5: yes\n"
        "Part B: Open-data format — 0.5\nScore 0.5: csv",
    ),
    "DPG 4": _q("DPG 4", "SQ", None),
}


def _agency(pid, kind):
    return Parastatal(id=pid, name=pid, type=kind)


UNITS = [
    Unit("Alpha", _agency("AWB", "water_supply_board")),
    Unit("Alpha", _agency("ATC", "transport_corporation")),
    Unit("Beta", _agency("BWB", "water_supply_board")),
]


@pytest.fixture
def draft(tmp_path):
    path = tmp_path / "templates" / "draft.xlsx"
    build_skeleton(QUESTIONS, UNITS, path, "PARA")
    return path


def test_skeleton_follows_the_template_conventions(draft):
    t = read_template(draft)
    assert t.problems == []
    assert set(t.questions) == {"DPG1a", "DPG1b", "DPG5a"}  # rollups and unscored left out
    assert t.labels == ["Alpha – AWB", "Alpha – ATC", "Beta – BWB"]
    assert [i.header for i in t.questions["DPG5a"].inputs] == [
        "Part A: Availability",
        "Part B: Open-data format",
    ]
    assert t.summary.overall_row and set(t.summary.question_rows) == {
        "DPG1",
        "DPG1a",
        "DPG1b",
        "DPG5a",
    }
    ws = load_workbook(draft)["DPG1b"]  # water boards only
    assert [ws.cell(row=r, column=2).value for r in (13, 14, 15)] == ["YES", "NO", "YES"]


def test_methodology_problems_are_reported():
    problems = check_questions(
        {**QUESTIONS, "DPG 9": _q("DPG 9", "SQ", 5, methodology="10 — mandatory")}
    )
    assert any("DPG 4" in p and "no Tag" in p for p in problems)
    assert any("DPG 9" in p and "up to 10" in p for p in problems)


def _plan(tmp_path, draft, rows):
    path = human.write_assignments_template(tmp_path / "assign.xlsx", {"PARA": draft.name})
    wb = load_workbook(path)
    for row in rows:
        wb["Assignments"].append(row)
    wb.save(path)
    register = {
        "alpha": CityConfig(slug="alpha", name="Alpha", state="S", ulg="U", aliases=["Alfa"])
    }
    return human.read_plan(path, draft.parent, register)


def test_assignments_split_a_question_between_interns(tmp_path, draft):
    plan, templates = _plan(
        tmp_path,
        draft,
        [
            ["PARA", "Asha", "DPG1*", "Alfa"],
            ["PARA", "Ravi", "DPG 1a", "Beta"],
            ["PARA", "Ravi", "DPG9", "All"],
        ],
    )
    asha, ravi = plan.assignments
    assert asha.questions == ["DPG1a", "DPG1b"] and asha.rows == ["Alpha – AWB", "Alpha – ATC"]
    assert ravi.rows == ["Beta – BWB"]
    assert any('no question "DPG9"' in p for p in plan.problems)


def test_merge_copies_typed_cells_and_flags_conflicts(tmp_path, draft):
    plan, templates = _plan(
        tmp_path, draft, [["PARA", "Asha", "DPG1a", "All"], ["PARA", "Ravi", "DPG1a", "Beta"]]
    )
    human.prepare_intern_copies(plan, templates, tmp_path)
    sheet = templates["PARA"].questions["DPG1a"]
    score = sheet.inputs[0].column
    for intern, value in [("Asha", 4), ("Ravi", 2)]:
        path = human.intern_file(tmp_path, "PARA", intern)
        wb = load_workbook(path)
        assert wb.sheetnames[0] == "Your assignment"
        assert wb["DPG5a"].sheet_state == "hidden"
        wb["DPG1a"].cell(row=sheet.rows["Beta – BWB"], column=score, value=value)
        wb.save(path)
    result = human.merge_vertical(plan, templates, "PARA", tmp_path)
    merged = load_workbook(result.path)["DPG1a"]
    assert merged.cell(row=sheet.rows["Beta – BWB"], column=score).value == 4
    assert str(
        merged.cell(row=sheet.rows["Beta – BWB"], column=sheet.points_column).value
    ).startswith("=")  # formulas are never overwritten
    assert any("Asha entered" in p and "Ravi entered" in p for p in result.problems)
    assert result.done["Asha"] == (1, 3)


def test_city_scores_average_agencies_and_results_formula(tmp_path, draft):
    t = read_template(draft)
    wb = load_workbook(draft)  # stand in for a calculated workbook: write values
    ws = wb[t.summary.sheet]
    for label, value in zip(t.labels, [8.0, 6.0, "NA"], strict=True):
        ws.cell(row=t.summary.overall_row, column=t.summary.labels[label], value=value)
        ws.cell(row=t.summary.overall_row + 1, column=t.summary.labels[label], value=0)
    wb.save(draft)
    scores = read_scores(t, draft, "PARA", "People")
    assert scores.city_scores() == {"Alpha": (7.0, False), "Beta": ("NA", False)}
    out = build_results([scores], tmp_path / "results.xlsx")
    sheet = load_workbook(out)["City scores"]
    assert sheet["B2"].value == 7.0 and sheet["C2"].value.startswith("=IFERROR(ROUND(AVERAGE(")


@pytest.mark.skipif(soffice() is None, reason="LibreOffice not installed")
def test_libreoffice_calculates_the_draft(tmp_path, draft):
    t = read_template(draft)
    wb = load_workbook(draft)
    for code, value in [("DPG1a", 4), ("DPG1b", 5)]:
        sheet = t.questions[code]
        for r in sheet.rows.values():
            wb[code].cell(row=r, column=sheet.inputs[0].column, value=value)
    for r in t.questions["DPG5a"].rows.values():
        for spec in t.questions["DPG5a"].inputs:
            wb["DPG5a"].cell(row=r, column=spec.column, value=0.5)
    wb.save(draft)
    recalculate(draft, tmp_path)
    scores = read_scores(t, draft, "PARA", "People")
    assert scores.questions[("DPG1", "Alpha – AWB")] == pytest.approx(9.0)
    assert scores.questions[("DPG1", "Alpha – ATC")] == pytest.approx(8.0)  # 1b is NA
    assert scores.overall["Alpha – AWB"] == pytest.approx(9.0)  # only DPG1 is an MQ here


def _step3(services, tmp_path, codes):
    import dataclasses

    from asics_agent.cities import load_register
    from asics_agent.question_bank import load_question_bank
    from asics_agent.scoring.ai import run_ai_scoring
    from asics_agent.scoring.skeleton import gather_units

    services.settings = dataclasses.replace(services.settings, scoring_dir=tmp_path / "scoring")
    register, _ = load_register(services.settings.city_register)
    questions, _ = load_question_bank(services.settings.question_bank)
    units, _ = gather_units(services.settings)
    draft = tmp_path / "scoring" / "templates" / "draft.xlsx"
    build_skeleton(questions, units, draft, "PARASTATAL")
    template = read_template(draft)
    report = run_ai_scoring(services, template, "PARASTATAL", register, codes=codes, workers=1)
    return template, report


def test_step3_needs_step2(services, run_steps, tmp_path):
    run_steps("sources")  # Step 1 only
    _, report = _step3(services, tmp_path, ["UPD1a"])
    assert report.cities == [] and report.scored == 0
    assert any("Step 2 (answer questions) hasn't been run" in p for p in report.problems)


def test_step3_scores_from_step2_citations_only(services, run_steps, tmp_path):
    from asics_agent.practice import SITE
    from asics_agent.scoring.ai import ai_file

    run_steps("sources")
    run_steps("answers")  # Step 2 for Sample City A, UPD 1 only
    template, report = _step3(services, tmp_path, ["UPD1a", "DPG1a"])
    assert report.cities == ["Sample City A"]
    assert any("Sample City B isn't ready" in p for p in report.problems)
    assert report.scored >= 1  # UPD 1a: Step 2 found evidence
    assert report.needs_person >= 1  # DPG 1a: Step 2 never answered it
    assert report.reasons.get("no_answer", 0) >= 1
    from asics_agent.scoring.ai import REVIEW_SHEET, _review_rows

    path = ai_file(tmp_path / "scoring", "PARASTATAL")
    book = load_workbook(path, data_only=True)
    assert book.sheetnames[0] == REVIEW_SHEET  # the one sheet to start from
    rows = _review_rows(book[REVIEW_SHEET])
    upd = rows[("UPD1a", "Sample City A – SWB")]
    assert upd["Status"] == "Scored by AI"
    assert upd["Step 2 status"] != "(no answer)" and upd["Step 2 answer"]  # what it worked from
    assert upd["Citation"].endswith("]") and upd["Link"] == SITE  # from the Citation Sheet
    if soffice():
        assert isinstance(upd["AI score"], (int, float))  # calculated by the workbook
    dpg = rows[("DPG1a", "Sample City A – SWB")]
    assert dpg["Status"] == "Left for a person" and dpg["Why left for a person"]
    ws = load_workbook(ai_file(tmp_path / "scoring", "PARASTATAL"))
    sheet = template.questions["UPD1a"]
    row = sheet.rows["Sample City A – SWB"]
    cell = lambda name: ws["UPD1a"].cell(row=row, column=sheet.evidence[name]).value  # noqa: E731
    assert cell("Link") == SITE  # copied from the Citation Sheet, not written by the AI
    assert cell("Act Name/Web name/Doc name").endswith("]")  # "<title> [<Citation ID>]"
    assert ws["UPD1a"].cell(row=row, column=sheet.inputs[0].column).value is not None
    dpg = template.questions["DPG1a"]
    comment = ws["DPG1a"].cell(row=dpg.rows["Sample City A – SWB"], column=dpg.comments_column)
    assert comment.value.startswith("AI: left for a person.")
    assert (
        ws["DPG1a"].cell(row=dpg.rows["Sample City A – SWB"], column=dpg.inputs[0].column).value
        is None
    )


def test_step3_discards_a_quote_not_in_the_source(services, run_steps, tmp_path, monkeypatch):
    from asics_agent.scoring import ai

    run_steps("sources")
    run_steps("answers")
    real = ai.call_json

    def made_up_quote(*args, **kwargs):
        reply, urls, notes = real(*args, **kwargs)
        return reply.model_copy(update={"quote": "words that are not in the source"}), urls, notes

    monkeypatch.setattr(ai, "call_json", made_up_quote)
    _, report = _step3(services, tmp_path, ["UPD1a"])
    assert report.scored == 0 and report.needs_person >= 1


def test_ai_copy_follows_a_changed_workbook_and_keeps_earlier_answers(tmp_path):
    import os

    from asics_agent.scoring.ai import ai_copy

    path = tmp_path / "draft.xlsx"
    build_skeleton(QUESTIONS, UNITS, path, "PARA")
    out = tmp_path / "ai" / "PARA_ai.xlsx"
    assert ai_copy(read_template(path), out) == ""  # first copy
    sheet = read_template(path).questions["DPG1a"]
    wb = load_workbook(out)
    wb["DPG1a"].cell(row=sheet.rows["Beta – BWB"], column=sheet.inputs[0].column, value=3)
    wb.save(out)
    more = {**QUESTIONS, "DPG 9": _q("DPG 9", "SQ", 5)}
    build_skeleton(more, UNITS, path, "PARA")  # the experts change the workbook
    os.utime(path, (out.stat().st_mtime + 5,) * 2)
    note = ai_copy(read_template(path), out)
    wb = load_workbook(out)
    assert "rebuilt" in note and "DPG9" in wb.sheetnames
    assert wb["DPG1a"].cell(row=sheet.rows["Beta – BWB"], column=sheet.inputs[0].column).value == 3


def test_a_different_question_under_the_same_code_is_not_scored():
    from asics_agent.scoring.ai import same_question

    assert same_question(
        "Is the parastatal mandated to prepare its sectoral plans in coordination with the ULG?",
        "Is the parastatal mandated to prepare its spatial and/or sectoral plans in "
        "coordination with the ULG?",
    )
    assert not same_question("Meeting minutes of governing board", "Annual Audit Report")


def test_your_score_is_kept_when_the_ai_scores_again(services, run_steps, tmp_path):
    from asics_agent.scoring.ai import REVIEW_SHEET, YOUR_SCORE, _review_rows, ai_file

    run_steps("sources")
    run_steps("answers")
    _step3(services, tmp_path, ["UPD1a"])
    path = ai_file(tmp_path / "scoring", "PARASTATAL")
    wb = load_workbook(path)
    ws = wb[REVIEW_SHEET]
    column = [c.value for c in ws[1]].index(YOUR_SCORE) + 1
    ws.cell(row=2, column=column, value=4)
    wb.save(path)
    _step3(services, tmp_path, ["UPD1a"])  # the AI scores the same rows again
    rows = _review_rows(load_workbook(path)[REVIEW_SHEET])
    assert [r[YOUR_SCORE] for r in rows.values()].count(4) == 1
