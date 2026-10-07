from pathlib import Path

from openpyxl import load_workbook

from asics_agent.practice import INDIA_CODE, MADE_UP, PLAN, SITE, TRANSPORT_SITE
from asics_agent.sources_workbook import CITATIONS_SHEET, PARASTATALS_SHEET


def test_step1_finds_parastatals_rates_websites_and_builds_citations(run_steps):
    result = run_steps("sources")
    parastatals = {p.id: p for p in result["parastatals"]}

    # The register listed only SWB; the agent found STC too.
    assert parastatals["SWB"].found_by == "Your team"
    assert parastatals["STC"].found_by == "Agent"
    # gov.in site linked from the state portal -> Official; .org site -> Probably official.
    assert parastatals["SWB"].website_check.band == "Official"
    assert parastatals["STC"].official_website == TRANSPORT_SITE
    assert parastatals["STC"].website_check.band == "Probably official"
    assert any("urban.example.gov.in" in r for r in parastatals["STC"].website_check.reasons)

    sources = {(s.parastatal_id, s.url): s for s in result["final_sources"]}
    assert ("SWB", MADE_UP) not in sources  # invented URL discarded
    india_code = sources[("SWB", INDIA_CODE)]
    assert india_code.verification_status == "Not verified" and not india_code.use_for_answers
    # The gap-filling pass found a source for UPD 1b.
    assert sources[("SWB", PLAN)].question_ids == ["UPD 1b"]
    assert all(s.citation_id for s in sources.values())

    wb = load_workbook(result["output_workbook"])
    assert {"How to Use", "Needs Attention", PARASTATALS_SHEET, CITATIONS_SHEET, "Coverage"} <= set(
        wb.sheetnames
    )


def test_step2_answers_only_from_the_citation_sheet(run_steps):
    run_steps("sources")
    result = run_steps("answers")
    answers = {(a.parastatal_id, a.question_id): a for a in result["answers"]}

    a = answers[("SWB", "UPD 1a")]
    assert a.status == "Answered — Verified Source"
    assert a.citation_id.startswith("SWB-") and a.citation_url == SITE
    assert a.proposed_score == 5 and "capped" in a.notes
    assert answers[("SWB", "UPD 1b")].citation_url == PLAN

    rollup = next(
        r for r in result["rollups"] if (r.parastatal_id, r.question_id) == ("SWB", "UPD 1")
    )
    assert rollup.proposed_score == 10
    assert not [
        i for i in result["final_issues"] if i.phase == "final_qa" and i.severity == "error"
    ]
    assert Path(result["output_workbook"]).name == "ASICS_Parastatal_Sample_City_A_Phase2.xlsx"


def _edit_sheet(path, sheet, header_name, match, column, value):
    wb = load_workbook(path)
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if match(row[header.index(header_name)]):
            row[header.index(column)].value = value
    wb.save(path)


def test_team_decisions_control_step2(run_steps):
    step1 = run_steps("sources")
    path = step1["output_workbook"]
    # The team turns off the plan document and excludes the transport corporation.
    _edit_sheet(
        path,
        CITATIONS_SHEET,
        "Official URL",
        lambda c: PLAN in str(c.value),
        "Use for Answers?",
        "No",
    )
    _edit_sheet(
        path, PARASTATALS_SHEET, "Parastatal ID", lambda c: c.value == "STC", "Include?", "No"
    )

    result = run_steps("answers")
    answers = {(a.parastatal_id, a.question_id): a for a in result["answers"]}
    assert not any(p == "STC" for p, _ in answers)
    unanswered = answers[("SWB", "UPD 1b")]
    assert unanswered.status == "Insufficient Evidence"
    assert unanswered.missing_evidence  # tells the team what kind of source to add


def test_team_added_citation_is_checked_and_used(run_steps):
    step1 = run_steps("sources")
    path = step1["output_workbook"]
    _edit_sheet(
        path,
        CITATIONS_SHEET,
        "Official URL",
        lambda c: PLAN in str(c.value),
        "Use for Answers?",
        "No",
    )
    wb = load_workbook(path)
    ws = wb[CITATIONS_SHEET]
    header = [c.value for c in ws[1]]
    row = [None] * len(header)
    for name, value in {
        "Parastatal": "SWB",
        "Source Title": "Plan (added by team)",
        "Official URL": PLAN,
        "Use for Answers?": "Yes",
    }.items():
        row[header.index(name)] = value
    ws.append(row)
    wb.save(path)

    result = run_steps("answers")
    answer = next(
        a for a in result["answers"] if (a.parastatal_id, a.question_id) == ("SWB", "UPD 1b")
    )
    assert answer.status == "Answered — Verified Source" and answer.citation_url == PLAN


def test_rerunning_step1_keeps_the_teams_work(run_steps):
    step1 = run_steps("sources")
    path = step1["output_workbook"]
    _edit_sheet(
        path,
        PARASTATALS_SHEET,
        "Parastatal ID",
        lambda c: c.value == "STC",
        "Your Notes",
        "Checked with the department",
    )
    ids_before = {s.url: s.citation_id for s in step1["final_sources"] if s.parastatal_id == "SWB"}

    again = run_steps("sources")
    stc = next(p for p in again["parastatals"] if p.id == "STC")
    assert stc.team_notes == "Checked with the department"
    ids_after = {s.url: s.citation_id for s in again["final_sources"] if s.parastatal_id == "SWB"}
    assert ids_after == ids_before  # Citation IDs are stable across runs
    assert list((Path(path).parent / "history").glob("*.xlsx"))  # previous version backed up


def test_step2_needs_step1_first(run_steps):
    result = run_steps("answers")
    assert any("Run Step 1" in i.message for i in result["issues"])


def test_question_tags_reach_sources_found_earlier_in_the_run(run_steps):
    result = run_steps("sources")
    site = next(s for s in result["final_sources"] if s.parastatal_id == "SWB" and s.url == SITE)
    assert "UPD 1a" in site.question_ids  # the official website, tagged by the later search


def test_insufficient_answers_carry_no_citation(run_steps, services):
    from asics_agent.practice import FakeClaude

    run_steps("sources")

    class Unhelpful(FakeClaude):
        def _generate(self, messages, stop=None, run_manager=None, **kwargs):
            if "TASK: answer_question" in messages[0].content:
                import json

                from langchain_core.messages import AIMessage
                from langchain_core.outputs import ChatGeneration, ChatResult

                reply = {
                    "sufficiency": "insufficient",
                    "citation_id": "SWB-01",
                    "notes": "The website doesn't cover this.",
                }
                message = AIMessage(
                    content=json.dumps(reply), response_metadata={"stop_reason": "end_turn"}
                )
                return ChatResult(generations=[ChatGeneration(message=message)])
            return super()._generate(messages, stop, run_manager, **kwargs)

    services.llm = Unhelpful()
    result = run_steps("answers")
    answer = next(a for a in result["answers"] if a.question_id == "UPD 1a")
    assert answer.status == "Insufficient Evidence"
    assert answer.citation_id == "" and answer.citation_url == ""
