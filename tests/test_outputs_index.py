"""outputs/START_HERE.xlsx: where each category and city is up to, with links."""

import dataclasses

from openpyxl import Workbook, load_workbook

from asics_agent.config import get_settings
from asics_agent.outputs_index import write_index


def test_index_lists_every_category_and_city_with_links(tmp_path):
    base = dataclasses.replace(
        get_settings(),
        outputs_dir=tmp_path,
        outputs_root=tmp_path,
        scoring_dir=tmp_path / "scoring",
    )
    from asics_agent.cities import load_register

    cities, _ = load_register(base.city_register)
    city = next(iter(cities.values())).name
    sources = tmp_path / "Parastatal" / "cities" / city / f"ASICS_{city}_Sources.xlsx"
    sources.parent.mkdir(parents=True)
    Workbook().save(sources)

    ws = load_workbook(write_index(base))["Start here"]
    rows = {(r[0], r[1]): r for r in ws.iter_rows(min_row=5, values_only=True)}
    assert rows[(city, "Parastatal")][2] == "Sources ready"
    assert rows[(city, "ULG")][2] == "Not started"
    link = next(c for c in ws[5 + list(rows).index((city, "Parastatal"))] if c.column == 5)
    assert link.hyperlink.target == f"Parastatal/cities/{city}/{sources.name}"
