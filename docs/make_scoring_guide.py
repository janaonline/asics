"""Writes docs/ASICS_Scoring_Guide.docx: the one-page guide for scoring experts.

    uv run --with python-docx python docs/make_scoring_guide.py
"""

import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("ASICS_Scoring_Guide.docx")
INK = RGBColor(0x11, 0x0F, 0x0F)

doc = Document()
sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.left_margin = sec.right_margin = Cm(1.6)
sec.top_margin = sec.bottom_margin = Cm(1.3)
normal = doc.styles["Normal"]
normal.font.name, normal.font.size = "Calibri", Pt(9)
normal.paragraph_format.space_after = Pt(2)


def add(p, text, bold=False, size=None):
    """Add text with **bold** markup."""
    for i, part in enumerate(text.split("**")):
        run = p.add_run(part)
        run.bold = bold or i % 2 == 1
        if size:
            run.font.size = Pt(size)
    return p


def para(text="", bold=False, size=None, after=2, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after, p.paragraph_format.space_before = Pt(after), Pt(before)
    return add(p, text, bold, size) if text else p


def heading(text):
    para(text, bold=True, size=11, before=5).runs[0].font.color.rgb = INK


def items(lines, numbered=False):
    for line in lines:
        p = doc.add_paragraph(style="List Number") if numbered else doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.left_indent = Cm(0.55)
        if not numbered:
            p.paragraph_format.first_line_indent = Cm(-0.35)
            p.add_run("•  ")
        add(p, line)


def table(rows, widths):
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style, t.alignment, t.autofit = "Table Grid", WD_TABLE_ALIGNMENT.CENTER, False
    for c, width in enumerate(widths):
        t.columns[c].width = Cm(width)
    for r, row in enumerate(rows):
        for c, text in enumerate(row):
            cell = t.cell(r, c)
            cell.width = Cm(widths[c])
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            add(cell.paragraphs[0], text, bold=r == 0, size=8.5)
            if r == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:fill"), "DDE7F0")
                cell._tc.get_or_add_tcPr().append(shading)


para("ASICS scoring: how experts score, and how scores add up", bold=True, size=15,
     after=1).runs[0].font.color.rgb = INK
para("Janaagraha · Annual Survey of India's City-Systems · for the research and scoring team",
     size=8.5, after=4)

heading("1. Where the scores come from")
para("Every vertical (Parastatal, UPD, and those added later) goes through the same three steps "
     "in the app: **Step 1** finds the official sources and builds the **Citation Sheet**, which "
     "the team reviews; **Step 2** answers each question using only that sheet, citing the source "
     "and quoting it; **Step 3** scores each answer in the experts' scoring workbook. The AI "
     "scores only when Step 2 found evidence, and only from the cited source; anything else is "
     "left for a person.")

heading("2. Scoring as an expert: one workbook, one sheet")
para("Open the vertical's workbook from the **Step 3 · Score** page (Open / Download), e.g. "
     "**PARASTATAL_ai.xlsx**, and work only on its first sheet, **Scores (start here)**: one row "
     "per question and agency (or per city, for city-level verticals such as UPD).")
items([
    "**Filter Status.** Left for a person (orange) rows have no AI score yet; Scored by AI rows "
    "have one to check.",
    "**Read the evidence**: Step 2 answer, Citation, Link (opens the source), Quote and the AI's "
    "reasoning. Why left for a person says what is missing.",
    "**Enter your score in ★ Your score (yellow)** in points, from 0 to the row's Max score, and "
    "a note in **★ Your comments**. Leave it blank where you agree with the AI.",
    "**Agrees with AI?** fills in by itself (red where you differ), and **Final score** shows the "
    "score that counts: yours if you entered one, otherwise the AI's.",
    "Save and close, keeping the file name, then press **Update scores** on the app's Scores page. "
    "Don't type in the other sheets: they are calculation sheets. Your entries are kept when the "
    "AI scores again.",
], numbered=True)
para("Your scores are also the **golden dataset**: kept next to the AI's own scores, they show "
     "where the AI is right and wrong, and are used to check the AI after every change.", after=3)

heading("3. How scores bubble up")
table([
    ["Level", "How it is calculated (by the experts' workbook formulas)"],
    ["Question (SQ)", "The **Final score**: your score where you entered one, otherwise the AI's "
     "points from its inputs. **NA** where it doesn't apply to that agency; **?** while not scored."],
    ["Main question (MQ)", "Sum of its applicable sub-questions, scaled to the MQ's maximum (usually "
     "10). NA sub-questions are left out; any ? makes the MQ ?."],
    ["Agency or city, per vertical", "Average of its scored main questions (out of 10). Provisional "
     "while any MQ is ? (shown with a count)."],
    ["City, per vertical", "**Parastatal:** average of the city's agencies (BWSSB, BMTC, BDA…). "
     "**City-level verticals (UPD…):** the city's own row."],
    ["City ASICS score", "Average of the city's vertical scores, equal weights."],
], [3.4, 14.4])
para("Every level updates as you type in Excel. Results are in **ASICS_2027_Scores.xlsx** and on "
     "the **Scores** page (after Update scores): each city and vertical, plus every question.",
     before=2)

heading("4. When a new vertical is added (e.g. Mobility)")
items([
    "A developer adds one settings file (**agent_setup/verticals/mobility.md**) naming its question "
    "bank, what it assesses (agencies or the city government), its agents and its scoring workbook. "
    "No new code.",
    "The experts provide the scoring workbook in the **same layout as UPD**: one sheet per question "
    "with its inputs and Points formula, a summary sheet, an overall-score row.",
    "It then appears in the app's **Vertical** picker and **Run steps**, runs Steps 1–3 the same "
    "way, gets its own **Scores (start here)** sheet, and adds a column to the results, so the "
    "city's ASICS score averages one more vertical.",
])

doc.save(OUT)
print(OUT)
