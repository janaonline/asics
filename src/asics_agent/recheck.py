"""Re-check the links in a finished workbook, e.g. just before a human review.

For every Citation Sheet link, and every answer's Citation URL, this re-opens the exact URL
with the same human-accessibility test used during the run. For answers it also confirms
the Evidence Excerpt is still on the page. Results go into new columns, in a copy of the
workbook saved next to the original (so "Open saved copy" links keep working).
"""

import tempfile
from collections import Counter
from datetime import date
from pathlib import Path

import httpx
from openpyxl import load_workbook

from asics_agent.links.accessibility import check_url
from asics_agent.links.text import _squash
from asics_agent.models import AccessCheck
from asics_agent.workbook import (
    _FILLS,
    CITATION_SHEET,
    _column_map,
    _norm,
    is_scoring_sheet,
)

RECHECK_COLUMNS = [("Link Re-checked On", []), ("Link Still Opens?", [])]
QUOTE_COLUMN = [("Quote Still on Page?", [])]


def _cell_url(cell) -> str:
    if cell.hyperlink and str(cell.hyperlink.target or "").startswith("http"):
        return cell.hyperlink.target
    value = str(cell.value or "").strip()
    return value if value.startswith("http") else ""


def _opens(check: AccessCheck) -> tuple[str, str]:
    reason = " ".join(check.reasons)
    if check.verification_status == "Verified":
        return "Yes", "green"
    if check.verification_status == "Human Verification Required":
        return f"Needs a person to open it: {reason}", "amber"
    return f"No: {reason}", "red"


def recheck_workbook(path: Path, http: httpx.Client, user_agent: str) -> tuple[Path, Counter]:
    wb = load_workbook(path)
    today = date.today().isoformat()
    results: dict[str, AccessCheck] = {}
    summary: Counter = Counter()

    with tempfile.TemporaryDirectory() as tmp:

        def check(url: str) -> AccessCheck:
            if url not in results:
                results[url] = check_url(url, http, Path(tmp), user_agent)
            return results[url]

        def text_of(result: AccessCheck) -> str:
            return (
                Path(result.content_path).read_text(encoding="utf-8") if result.content_path else ""
            )

        sheets = [
            ws
            for ws in wb.worksheets
            if _norm(ws.title) == _norm(CITATION_SHEET) or is_scoring_sheet(ws.title)
        ]
        for ws in sheets:
            is_scoring = is_scoring_sheet(ws.title)
            headers = {_norm(c.value): c.column for c in ws[1] if c.value}
            url_column = headers.get(_norm("Citation URL" if is_scoring else "Official URL"))
            if url_column is None:
                continue
            mapping = _column_map(ws, RECHECK_COLUMNS + (QUOTE_COLUMN if is_scoring else []))
            quote_column = headers.get(_norm("Evidence Excerpt"))
            for r in range(2, ws.max_row + 1):
                url = _cell_url(ws.cell(row=r, column=url_column))
                if not url:
                    continue
                result = check(url)
                opens, colour = _opens(result)
                summary["links opening" if opens == "Yes" else "links with problems"] += 1
                ws.cell(row=r, column=mapping["Link Re-checked On"], value=today)
                cell = ws.cell(row=r, column=mapping["Link Still Opens?"], value=opens)
                cell.fill = _FILLS[colour]
                if is_scoring and quote_column:
                    quote = str(ws.cell(row=r, column=quote_column).value or "")
                    if not quote:
                        continue
                    if opens != "Yes":
                        found, colour = "Could not check: the link didn't open", "amber"
                    elif _squash(quote) in _squash(text_of(result)):
                        found, colour = "Yes", "green"
                    else:
                        found, colour = "No: the page has changed. Use the saved copy", "red"
                        summary["quotes no longer on the page"] += 1
                    cell = ws.cell(row=r, column=mapping["Quote Still on Page?"], value=found)
                    cell.fill = _FILLS[colour]

    out = path.with_name(f"{path.stem} (re-checked {today}){path.suffix}")
    wb.save(out)
    return out, summary


def summary_text(summary: Counter) -> str:
    opening = summary.get("links opening", 0)
    broken = summary.get("links with problems", 0)
    moved = summary.get("quotes no longer on the page", 0)
    if not opening and not broken:
        return "No links were found in this workbook."
    parts = [f"{opening} link{'s' if opening != 1 else ''} still open properly."]
    if broken:
        parts.append(
            f"{broken} link{'s have' if broken != 1 else ' has'} a problem: see the "
            '"Link Still Opens?" column.'
        )
    if moved:
        parts.append(
            f"{moved} quote{'s are' if moved != 1 else ' is'} no longer on the page: "
            'see "Quote Still on Page?" and use the saved copy.'
        )
    return " ".join(parts)
