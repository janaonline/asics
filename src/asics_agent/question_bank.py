"""Loads the ASICS parastatal question bank workbook.

The bank's IDs are written inconsistently ("DPG-2b", "DPG5a", "DPG d"), so IDs are
normalised to "<PILLAR> <number><letter>" and every fix is reported as an Issue.
"""

import re
from pathlib import Path

from openpyxl import load_workbook

from asics_agent.models import Issue, Question

SHEET = "Question Bank"
REQUIRED_COLUMNS = [
    "Sl. No.",
    "City-Systems Pillar",
    "Question",
    "Tag",
    "Score / Max Score",
    "Assessment Level",
    "Applicability",
    "Detailed Methodology",
    "Evidence / Source Requirement",
]

_ID = re.compile(r"^\s*([A-Za-z]+)[\s\-]*(\d*)\s*([a-z]?)\s*$")
_SECTION = re.compile(r"\(([A-Z]+)\)")


def _clean(value) -> str:
    return "" if value is None else str(value).strip()


def _score(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


# Columns a vertical's bank may leave out: every question then applies to every unit.
OPTIONAL_COLUMNS = {"Sl. No.": "", "Applicability": "Common.", "Evidence / Source Requirement": ""}
_PART = re.compile(r"^\s*([A-Za-z]+\s*\d+[a-z])\s*(\d+)\s*$")  # e.g. UPD23a1: a part of UPD23a


def load_question_bank(
    path: Path, sheet: str | None = None, columns: dict[str, str] | None = None
) -> tuple[dict[str, Question], list[Issue]]:
    """Load a vertical's question bank.

    `sheet` names the sheet (default: "Question Bank", else the first sheet with the
    columns); `columns` maps our column names to the bank's own, e.g.
    {"Tag": "MQ_SQ", "Score / Max Score": "Max Score"} (set in the vertical's settings).
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    rename = {actual.strip(): ours for ours, actual in (columns or {}).items()}
    wanted = set(REQUIRED_COLUMNS) - set(OPTIONAL_COLUMNS)

    def headers(w) -> set[str]:
        first = next(w.iter_rows(max_row=1, values_only=True), ())
        return {rename.get(_clean(c), _clean(c)) for c in first}

    ws = (
        wb[sheet]
        if sheet and sheet in wb.sheetnames
        else wb[SHEET]
        if SHEET in wb.sheetnames
        else next(
            (
                w
                for w in wb.worksheets  # e.g. the scoring team's "Parastatal Methodology" sheet
                if headers(w) >= wanted
            ),
            wb.worksheets[0],
        )
    )
    rows = list(ws.iter_rows(values_only=True))
    header = [rename.get(_clean(h), _clean(h)) for h in rows[0]]
    issues: list[Issue] = []
    missing = [c for c in REQUIRED_COLUMNS if c not in header and c not in OPTIONAL_COLUMNS]
    if missing:
        issues.append(
            Issue(
                phase="initial_checks",
                severity="error",
                audience="team",
                message="The Question Bank sheet is missing these columns: "
                f"{', '.join(missing)}. Add them back (with the same names) and try again.",
            )
        )
        return {}, issues
    col = {name: header.index(name) for name in header if name}
    defaults = {name: value for name, value in OPTIONAL_COLUMNS.items() if name not in col}

    def cell(row, name: str) -> str:
        if name in defaults:
            return defaults[name]
        i = col[name]
        return _clean(row[i]) if i < len(row) else ""

    questions: dict[str, Question] = {}
    pillar_code, pillar_name, parent = "", "", None
    for row_no, row in enumerate(rows[1:], start=2):
        raw_id = cell(row, "City-Systems Pillar")
        if not raw_id:
            # Section header row, e.g. "Urban Planning & Design (UPD)"
            pillar_name = _clean(row[0])
            if not pillar_name:
                continue
            match = _SECTION.search(pillar_name)
            pillar_code = match.group(1) if match else ""
            parent = None
            continue

        if part := _PART.match(raw_id):  # scored as part of its question's own sheet
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="info",
                    ref=raw_id,
                    message=f"Row {row_no}: {raw_id} is a part of {part.group(1)}; it is "
                    "answered with that question.",
                )
            )
            continue
        match = _ID.match(raw_id)
        if not match:
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="error",
                    ref=raw_id,
                    audience="team",
                    message=f'Question Bank row {row_no}: the ID "{raw_id}" isn\'t in a form '
                    'we understand (e.g. "UPD 1a"). Please correct it.',
                )
            )
            continue
        prefix, number, letter = match.groups()
        prefix = prefix.upper()
        tag = cell(row, "Tag").upper()
        if not pillar_code:
            pillar_code = prefix
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="warning",
                    ref=raw_id,
                    message=f"Section {pillar_name!r} has no (CODE); using {prefix}",
                )
            )
        if prefix != pillar_code:
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="warning",
                    ref=raw_id,
                    message=f"Row {row_no}: ID prefix {prefix} differs from "
                    f"section code {pillar_code}",
                )
            )
        if not number and tag == "SQ" and parent:
            number = re.sub(r"\D", "", parent.id.split(" ")[1])
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="warning",
                    ref=raw_id,
                    audience="team",
                    message=f'Question Bank row {row_no}: the ID "{raw_id}" has no number. '
                    f"We treated it as part of {parent.id}. Please correct the ID in the "
                    "question bank.",
                )
            )
        qid = f"{pillar_code} {number}{letter}"
        if qid != raw_id:
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="info",
                    ref=qid,
                    message=f"Row {row_no}: normalised ID {raw_id!r} -> {qid!r}",
                )
            )
        if qid in questions:
            issues.append(
                Issue(
                    phase="initial_checks",
                    severity="error",
                    ref=qid,
                    audience="team",
                    message=f'Question Bank row {row_no}: the ID "{raw_id}" is used twice. '
                    "Each question needs its own ID; the second one was skipped.",
                )
            )
            continue

        question = Question(
            id=qid,
            raw_id=raw_id,
            row=row_no,
            pillar=pillar_code,
            pillar_name=pillar_name,
            text=cell(row, "Question"),
            tag="MQ" if tag == "MQ" else "SQ",
            max_score=_score(cell(row, "Score / Max Score")),
            assessment_level=cell(row, "Assessment Level") or None,
            unit=_clean(row[col.get("Unit of Assessment", -1)])
            if "Unit of Assessment" in col
            else None,
            applicability=cell(row, "Applicability"),
            methodology=cell(row, "Detailed Methodology"),
            evidence_requirement=cell(row, "Evidence / Source Requirement"),
            rationale=cell(row, "Rationale") if "Rationale" in col else "",
        )
        if question.tag == "MQ":
            parent = question
        elif parent is not None:
            question.parent_id = parent.id
            parent.children.append(qid)
        questions[qid] = question
    return questions, issues
