"""outputs/START_HERE.xlsx: the one file that says where everything is.

One row per city and category: how far it has got, its latest workbooks (click to open), and
what to do next. A second sheet explains every folder. It is rewritten after every run (and
when the app starts), from the files on disk only, so it is quick and always current.

The outputs folder:

  outputs/
    START_HERE.xlsx
    <Category>/cities/<City>/   that category's work for that city (see FOLDERS)
    <Category>/batches/         one index workbook per multi-city run
    <Category>/practice/        practice runs (sample data)
    _app/                       the app's own files (error log, uploads)
    _old/                       earlier layouts, kept for reference
"""

import os
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

INDEX = "START_HERE.xlsx"
APP_DIR = "_app"
FOLDERS = [
    ("START_HERE.xlsx", "This file. Start here."),
    (
        "<Category>/cities/<City>/ASICS_<City>_Sources.xlsx",
        "Step 1: the parastatals (or the ULG) and the Citation Sheet. Your team reviews and edits "
        "this one.",
    ),
    (
        "<Category>/cities/<City>/answers/<run>/ASICS_<Category>_<City>_Phase2.xlsx",
        "Step 2: the answers, one folder per run. The newest is the one to use; older runs are "
        "kept for comparison.",
    ),
    (
        "scoring/2027/ai/<CODE>_ai.xlsx (outside outputs)",
        "Step 3: the scoring workbook. Its first sheet, Scores (start here), is the one to use.",
    ),
    (
        "<Category>/cities/<City>/history/",
        "Earlier versions of the Sources workbook, saved before each Step 1 run replaced it.",
    ),
    ("<Category>/cities/<City>/checks/", "Reports from free check-only runs."),
    (
        "<Category>/cities/<City>/evidence/",
        "Saved copies of every source as it was when checked (the 'Open saved copy' links). "
        "Needed by Steps 2 and 3: don't delete.",
    ),
    (
        "<Category>/cities/<City>/links.db, sources_summary.json, answers/<run>/links.db, "
        "run_summary.json",
        "Technical records of each run, for developers.",
    ),
    ("<Category>/batches/<run>/", "For runs over several cities: one index workbook listing them."),
    ("<Category>/practice/", "Practice runs with sample data (not real research)."),
    ("_app/", "The app's own files: error log, uploaded workbooks, last city and category chosen."),
    ("_old/", "Files from earlier versions of the tool, kept for reference."),
]
NEXT = {
    "Not started": "Run Step 1 (Find sources)",
    "Sources ready": "Review the Sources workbook, then run Step 2",
    "Sources changed": "Sources were edited after the answers: run Step 2 again",
    "Answered": "Review the answers, then run Step 3 (Score)",
    "Scored": "Check the scores in the scoring workbook (yellow column for your own)",
}
_HEAD = PatternFill("solid", fgColor="DDE7F0")


def app_dir(settings) -> Path:
    """outputs/_app: the app's own files, whatever category is chosen."""
    root = getattr(settings, "outputs_root", None) or settings.outputs_dir
    path = root / APP_DIR
    path.mkdir(parents=True, exist_ok=True)
    return path


def _latest(paths) -> Path | None:
    paths = [p for p in paths if "re-checked" not in p.name]
    return max(paths, key=lambda p: p.stat().st_mtime, default=None)


def city_row(vs, city_name: str, scored: dict) -> dict:
    """Where one city is up to in one category, from the files on disk."""
    from asics_agent.sources_workbook import sources_path, workspace_for
    from asics_agent.workbook import answers_files

    folder = workspace_for(vs.outputs_dir, city_name)
    sources = sources_path(folder, city_name)
    answers = _latest(answers_files(folder, vs))
    if not sources.exists():
        stage = "Not started"
    elif answers is None:
        stage = "Sources ready"
    elif answers.stat().st_mtime < sources.stat().st_mtime:
        stage = "Sources changed"
    else:
        stage = "Answered"
        when = scored.get(city_name)
        if when and datetime.fromisoformat(when).timestamp() >= answers.stat().st_mtime:
            stage = "Scored"
    return {
        "folder": folder if folder.exists() else None,
        "stage": stage,
        "sources": sources if sources.exists() else None,
        "answers": answers,
    }


def write_index(settings) -> Path:
    """Write outputs/START_HERE.xlsx for every category and every city in the register."""
    from asics_agent.categories import category_settings, load_categories
    from asics_agent.cities import load_register
    from asics_agent.scoring.ai import ai_file, scored_cities

    root = getattr(settings, "outputs_root", None) or settings.outputs_dir
    root.mkdir(parents=True, exist_ok=True)
    path = root / INDEX
    found, _ = load_categories(settings.agent_setup_dir)
    register, _ = load_register(settings.city_register)

    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    ws.append([f"ASICS outputs · updated {datetime.now():%d %b %Y %H:%M}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(
        [
            "One row per city and category. Click a file name to open it. The Step 3 "
            "scoring workbook covers every city of its category."
        ]
    )
    ws.append([])
    columns = [
        "City",
        "Category",
        "Where it's up to",
        "Next step",
        "1 · Sources workbook (review)",
        "2 · Answers workbook (latest)",
        "3 · Scoring workbook",
        "City folder",
    ]
    ws.append(columns)
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True), _HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for category in found.values():
        vs = category_settings(settings, category.name)
        scored = scored_cities(vs.scoring_dir, vs.scoring_code)
        scoring = ai_file(vs.scoring_dir, vs.scoring_code)
        for city in register.values():
            info = city_row(vs, city.name, scored)
            ws.append([city.name, category.title, info["stage"], NEXT[info["stage"]]])
            r = ws.max_row
            links = [
                info["sources"],
                info["answers"],
                scoring if scoring.exists() and info["stage"] == "Scored" else None,
                info["folder"],
            ]
            for c, target in enumerate(links, start=5):
                if target is None:
                    ws.cell(row=r, column=c, value="–")
                    continue
                cell = ws.cell(row=r, column=c, value=target.name)
                cell.hyperlink = os.path.relpath(target, root)
                cell.font = Font(color="1F4E9C", underline="single")
    for letter, width in zip("ABCDEFGH", [14, 12, 16, 44, 34, 40, 22, 14], strict=True):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "C5"

    guide = wb.create_sheet("What's in each folder")
    guide.append(["File or folder", "What it is"])
    for cell in guide[1]:
        cell.font, cell.fill = Font(bold=True), _HEAD
    for row in FOLDERS:
        guide.append(list(row))
    guide.column_dimensions["A"].width = 70
    guide.column_dimensions["B"].width = 90
    for row in guide.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    partial = path.with_name(f".{path.name}.writing")
    wb.save(partial)
    partial.replace(path)
    return path


def refresh_index(settings) -> None:
    """Rewrite the index, never stopping a run if it can't (e.g. the file is open in Excel)."""
    import logging

    if "practice" in settings.outputs_dir.parts:  # practice runs never touch the real index
        return
    try:
        write_index(settings)
    except Exception:  # noqa: BLE001 (the index is a convenience)
        logging.exception("couldn't update outputs/START_HERE.xlsx")
