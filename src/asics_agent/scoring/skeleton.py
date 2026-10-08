"""Builds a draft scoring workbook for a vertical from its question/methodology sheet.

The draft follows the same conventions as the expert team's UPD workbook (see template.py),
so the same tools fill, merge and read it. It is a starting point for the experts: each
question gets one simple input ("Score", or one per "Part A/B…" when the methodology splits
the points into parts), and the experts replace these with their own inputs and formulas.

Rows are one per city and agency ("Bengaluru – BWSSB"); "Applies?" comes from the question's
Applicability and the agency's type. The summary rolls sub-questions up to main questions and
main questions up to an overall score per agency, like UPD_SUMMARY_RAW; a city sheet averages
each city's agencies.
"""

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from asics_agent.applicability import match_rule
from asics_agent.models import Issue, Parastatal, Question
from asics_agent.scoring.template import COMMENTS, EVIDENCE, POINTS

HEADER_ROW = 12
FIRST_ROW = HEADER_ROW + 1
_HEAD = PatternFill("solid", fgColor="DDE7F0")
_INPUT = PatternFill("solid", fgColor="FFF6D5")
_GREY = PatternFill("solid", fgColor="E7E6E6")
_PART = re.compile(r"^\s*Part\s+([A-Z])\s*:\s*(.+?)\s*[—–-]+\s*([\d.]+)\s*$", re.M)
_LEVEL = re.compile(r"(?:^|\n)\s*(?:Score\s+)?(\d+(?:\.\d+)?)\s*(?:[—–-]|:|if\b)", re.I)
DETAILS = [  # label on the question sheet -> column of the Questions sheet (1-based)
    ("Group", 3),
    ("Original ID", 2),
    ("MQ_SQ", 4),
    ("Question", 5),
    ("Max Score", 6),
    ("Assessment level", 7),
    ("Applicability", 8),
    ("Detailed Methodology", 9),
    ("Notes for scorers", 10),
]
QUESTION_COLUMNS = [
    "Q No",
    "Original ID",
    "Group",
    "MQ_SQ",
    "Question",
    "Max Score",
    "Assessment level",
    "Applicability",
    "Detailed Methodology",
    "Notes for scorers",
]


@dataclass
class Unit:
    """One row of every question sheet: an agency in a city."""

    city: str
    agency: Parastatal

    @property
    def label(self) -> str:
        return f"{self.city} – {self.agency.id}"


def sheet_code(question_id: str) -> str:
    """'DPG 5a' -> 'DPG5a' (sheet names and summary codes have no spaces, as in UPD)."""
    return question_id.replace(" ", "")


def _parts(methodology: str, max_score: float) -> list[tuple[str, float]]:
    """[("Part A: Availability and currency", 0.5), …] when the parts add up to the max."""
    parts = [(f"Part {m[1]}: {m[2].strip()}", float(m[3])) for m in _PART.finditer(methodology)]
    if len(parts) >= 2 and abs(sum(p for _, p in parts) - max_score) < 1e-9:
        return parts
    return []


def _number(value: float) -> str:
    return f"{value:g}"


def check_questions(questions: dict[str, Question]) -> list[str]:
    """Plain-language problems the experts should look at before scoring starts."""
    problems = []
    for q in questions.values():
        if q.max_score is None:
            problems.append(
                f"{q.raw_id} has no Tag or Max Score, so it is left out of the "
                "scores. Add them if it should be scored."
            )
            continue
        if q.tag == "MQ" and q.children:
            kids = [questions[c] for c in q.children if questions[c].max_score is not None]
            total = sum(k.max_score for k in kids)
            if kids and abs(total - q.max_score) > 1e-9:
                problems.append(
                    f"{q.raw_id} is worth {_number(q.max_score)} but its "
                    f"sub-questions add up to {_number(total)}. The summary scales "
                    "them to the main question's max score, as in UPD."
                )
        if match_rule(q.applicability) is None:
            problems.append(
                f'{q.raw_id}: the Applicability "{q.applicability}" isn\'t one we '
                "recognise, so the draft assumes it applies to every agency."
            )
        if not q.children:
            levels = [float(v) for v in _LEVEL.findall(q.methodology)]
            if levels and max(levels) > q.max_score:
                problems.append(
                    f"{q.raw_id}: the methodology gives scores up to "
                    f"{_number(max(levels))}, but the Max Score is "
                    f"{_number(q.max_score)}. Say how the score should be scaled."
                )
    return problems


def _scored(questions: dict[str, Question]) -> dict[str, Question]:
    """Questions that get a score: those with a max score. Unscored ones are detached from
    their main question (a sub-question without a score must not count towards it)."""
    scored = {k: q.model_copy(deep=True) for k, q in questions.items() if q.max_score is not None}
    for q in scored.values():
        q.children = [c for c in q.children if c in scored]
    return scored


def _applies(question: Question, agency: Parastatal) -> bool:
    rule = match_rule(question.applicability)
    return True if rule is None else rule[1](agency)


def _header(ws, row: int, values: list, widths: dict[int, int] | None = None):
    for c, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=c, value=value)
        cell.font, cell.fill = Font(bold=True), _HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for c, w in (widths or {}).items():
        ws.column_dimensions[get_column_letter(c)].width = w


def _questions_sheet(wb, title: str, questions: dict[str, Question]):
    ws = wb.create_sheet(title)
    _header(ws, 1, QUESTION_COLUMNS, {1: 10, 2: 10, 3: 30, 5: 60, 8: 30, 9: 80, 10: 50})
    for q in questions.values():
        ws.append(
            [
                sheet_code(q.id),
                q.raw_id,
                q.pillar_name,
                q.tag if q.max_score else "",
                q.text,
                q.max_score,
                q.assessment_level,
                q.applicability,
                q.methodology,
                q.evidence_requirement,
            ]
        )
    ws.freeze_panes = "B2"
    return ws


def _question_sheet(wb, q: Question, units: list[Unit], questions_sheet: str):
    code = sheet_code(q.id)
    ws = wb.create_sheet(code)
    ws["A1"], ws["B1"] = "Question Code", code
    ws["C1"] = "<Sheet name and this code must match>"
    lookup = f"'{questions_sheet}'!$A:$J"
    for r, (label, column) in enumerate(DETAILS, start=2):
        ws.cell(row=r, column=1, value=label).font = Font(bold=True)
        ws.cell(row=r, column=2, value=f'=VLOOKUP($B$1,{lookup},{column},0)&""')
    max_cell = "$B$6"  # Max Score, as a number
    ws["B6"] = f"=VLOOKUP($B$1,{lookup},6,0)"
    for r in range(2, 11):
        ws.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    parts = _parts(q.methodology, q.max_score)
    inputs = [
        (name, f"SELECT {_number(points)} / 0", [_number(points), "0"]) for name, points in parts
    ] or [("Score", f"ENTER A NUMBER 0 TO {_number(q.max_score)}", None)]
    headers = ["City", "Applies?", *EVIDENCE, *[name for name, _, _ in inputs], POINTS, COMMENTS]
    marks = [
        "",
        "PRE-FILLED - DO NOT ALTER",
        *[""] * len(EVIDENCE),
        *[mark for _, mark, _ in inputs],
        "DO NOT ALTER FORMULA",
        "",
    ]
    for c, mark in enumerate(marks, start=1):
        if mark:
            ws.cell(row=HEADER_ROW - 1, column=c, value=mark).font = Font(bold=True, size=9)
    _header(ws, HEADER_ROW, headers, {1: 28, 2: 10, 3: 30, 8: 10, 9: 40})
    first_input = 3 + len(EVIDENCE)
    input_cols = list(range(first_input, first_input + len(inputs)))
    points_col = input_cols[-1] + 1
    last = FIRST_ROW + len(units) - 1
    for i, unit in enumerate(units):
        r = FIRST_ROW + i
        applies = _applies(q, unit.agency)
        ws.cell(row=r, column=1, value=unit.label)
        ws.cell(row=r, column=2, value="YES" if applies else "NO")
        cells = [f"{get_column_letter(c)}{r}" for c in input_cols]
        if parts:
            filled = f"COUNTA({cells[0]}:{cells[-1]})"
            value = f'IF({filled}<{len(cells)},"",SUM({cells[0]}:{cells[-1]}))'
        else:
            value = f'IF({cells[0]}="","",MIN({cells[0]},{max_cell}))'
        ws.cell(row=r, column=points_col, value=f'=IF($B{r}="NO","NA",{value})')
        for c in input_cols:
            ws.cell(row=r, column=c).fill = _GREY if not applies else _INPUT
        if not applies:
            for c in range(1, points_col + 2):
                if c not in input_cols:
                    ws.cell(row=r, column=c).fill = _GREY
    for c, (_, _, options) in zip(input_cols, inputs, strict=True):
        letter = get_column_letter(c)
        if options:
            dv = DataValidation(type="list", formula1=f'"{",".join(options)}"', allow_blank=True)
        else:
            dv = DataValidation(
                type="decimal",
                operator="between",
                formula1="0",
                formula2=_number(q.max_score),
                allow_blank=True,
            )
        dv.error = (
            f"See the instruction above the '{ws.cell(row=HEADER_ROW, column=c).value}' column."
        )
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f"{letter}{FIRST_ROW}:{letter}{last}")
    ws.freeze_panes = ws.cell(row=FIRST_ROW, column=2)
    return ws


def _leaf_formula(row: int, col: str, last_row: int) -> str:
    """The UPD_SUMMARY_RAW leaf formula: the question sheet's Points for this column's label,
    "?" while blank, "ER" when the sheet's code does not match."""
    sheet = f'INDIRECT("\'"&$A{row}&"\'!$A${HEADER_ROW}:$Z${last_row}")'
    header = f'INDIRECT("\'"&$A{row}&"\'!$A${HEADER_ROW}:$Z${HEADER_ROW}")'
    points = f'VLOOKUP({col}$1,{sheet},MATCH("{POINTS}",{header},0),0)'
    return (
        f'=IF($F{row}=0,"NA",IF($A{row}=INDIRECT("\'"&$A{row}&"\'!$B$1"),'
        f'IF({points}="","?",{points}),"ER"))'
    )


def _rollup_formula(row: int, col: str, kids: list[int]) -> str:
    """The UPD main-question formula: sum of sub-questions over the max of those that apply,
    scaled to the main question's max."""
    cells = [f"{col}{k}" for k in kids]
    any_of = lambda v: ",".join(f'{c}="{v}"' for c in cells)  # noqa: E731
    denominator = "+".join(f'({c}<>"NA")*$F${k}' for c, k in zip(cells, kids, strict=True))
    return (
        f'=IF(OR({any_of("ER")}),"ER",IF(OR({any_of("?")}),"?",'
        f'IF(AND({any_of("NA")}),"NA",SUM({",".join(cells)})/({denominator})*$F{row})))'
    )


def _summary_sheet(
    wb, vertical: str, questions: dict[str, Question], units: list[Unit], questions_sheet: str
):
    ws = wb.create_sheet(f"{vertical}_SUMMARY_RAW", 1)
    labels = [u.label for u in units]
    _header(
        ws,
        1,
        [
            "ASICS 2027_Report_Q No",
            "Group",
            "MQ_SQ",
            "Question",
            "Assessment level",
            "Max Score",
            *labels,
        ],
        {1: 14, 2: 24, 4: 50, 5: 18},
    )
    ws.append(["ENTER", *["Derived"] * (5 + len(labels))])
    rows = {code: 3 + i for i, code in enumerate(questions)}
    last_row = FIRST_ROW + len(units) - 1
    lookup = f"'{questions_sheet}'!$A:$J"
    for qid, q in questions.items():
        r = rows[qid]
        ws.cell(row=r, column=1, value=sheet_code(qid))
        for c, column in [(2, 3), (3, 4), (4, 5), (5, 7)]:
            ws.cell(row=r, column=c, value=f'=VLOOKUP($A{r},{lookup},{column},0)&""')
        ws.cell(row=r, column=6, value=f"=VLOOKUP($A{r},{lookup},6,0)")
        for i in range(len(labels)):
            col = get_column_letter(7 + i)
            ws.cell(
                row=r,
                column=7 + i,
                value=_rollup_formula(r, col, [rows[c] for c in q.children])
                if q.children
                else _leaf_formula(r, col, last_row),
            )
        if q.tag == "MQ":
            for c in range(1, 7 + len(labels)):
                ws.cell(row=r, column=c).font = Font(bold=True)
    end = 2 + len(questions)
    overall = end + 3
    ws.cell(row=overall - 1, column=4, value=f"{vertical}_Scores").font = Font(bold=True)
    texts = [
        f"{vertical}_OVERALL_SCORE",
        'No. of MQs not yet scored ("?") - overall score is provisional while > 0',
        "No. of MQs marked NA (excluded from overall score)",
    ]
    for k, text in enumerate(texts):
        ws.cell(row=overall + k, column=4, value=text).font = Font(bold=True)
    for i, label in enumerate(labels):
        c = get_column_letter(7 + i)
        ws.cell(row=overall - 1, column=7 + i, value=label).font = Font(bold=True)
        mq = f'$C$3:$C${end},"MQ"'
        ws.cell(
            row=overall,
            column=7 + i,
            value=(
                f"=IFERROR(SUMIFS({c}$3:{c}${end},{mq})/(COUNTIFS({mq})-COUNTIFS({mq},"
                f'{c}$3:{c}${end},"NA")),"NA")'
            ),
        )
        ws.cell(row=overall + 1, column=7 + i, value=f'=COUNTIFS({mq},{c}$3:{c}${end},"?")')
        ws.cell(row=overall + 2, column=7 + i, value=f'=COUNTIFS({mq},{c}$3:{c}${end},"NA")')
    ws.freeze_panes = "G3"
    return ws, overall


def _city_sheet(wb, vertical: str, units: list[Unit], summary: str, overall: int):
    """City score = average of the city's agencies' overall scores (equal weights)."""
    ws = wb.create_sheet(f"{vertical}_CITY_SCORES", 2)
    _header(
        ws,
        1,
        ["City", "Agencies", f"{vertical} score", "Agencies with a score"],
        {1: 22, 2: 50, 3: 16, 4: 20},
    )
    end = get_column_letter(6 + len(units))
    labels = f"'{summary}'!$G${overall - 1}:${end}${overall - 1}"
    scores = f"'{summary}'!$G${overall}:${end}${overall}"
    cities = list(dict.fromkeys(u.city for u in units))
    for r, city in enumerate(cities, start=2):
        agencies = ", ".join(u.agency.id for u in units if u.city == city)
        ws.append(
            [
                city,
                agencies,
                f'=IFERROR(AVERAGEIFS({scores},{labels},$A{r}&" – *"),"NA")',
                f'=COUNTIFS({labels},$A{r}&" – *",{scores},">=0")',
            ]
        )
    ws.cell(
        row=len(cities) + 3,
        column=1,
        value=(
            "A city's score is the average of its agencies' overall scores. Agencies with no "
            "applicable questions (NA) are left out."
        ),
    ).font = Font(italic=True)


def build_skeleton(
    questions: dict[str, Question],
    units: list[Unit],
    path: Path,
    vertical: str = "PARASTATAL",
    issues: list[Issue] | None = None,
) -> list[str]:
    """Write the draft workbook; return the problems listed on its "Check before use" sheet."""
    problems = [i.message for i in issues or [] if i.severity != "info"]
    problems += check_questions(questions)
    scored = _scored(questions)
    wb = Workbook()
    guide = wb.active
    guide.title = "Check before use"
    guide.column_dimensions["A"].width = 120
    guide.append([f"DRAFT {vertical} scoring workbook - generated from the methodology sheet"])
    guide["A1"].font = Font(bold=True, size=14)
    for line in [
        "",
        "This draft follows the same layout as the UPD scoring workbook, so the app can fill, "
        "merge and read it. The scoring experts should:",
        "  1. Replace each question sheet's simple inputs (yellow) and Points formula with the "
        "real ones, keeping the row-11 instructions (e.g. SELECT YES / NO) and the column names "
        f"'{POINTS}' and '{COMMENTS}'.",
        "  2. Check 'Applies?' (filled in from each agency's type and the question's "
        "Applicability). To change it, fix the agency in the city register or Sources workbook "
        "and generate again.",
        "  3. Fix the points below in the methodology sheet and generate again.",
        "",
        "Points to check:",
        *([f"  • {p}" for p in problems] or ["  (none)"]),
    ]:
        guide.append([line])
    for row in guide.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    questions_sheet = f"{vertical} Questions"
    _questions_sheet(wb, questions_sheet, questions)
    summary, overall = _summary_sheet(wb, vertical, scored, units, questions_sheet)
    _city_sheet(wb, vertical, units, summary.title, overall)
    for q in scored.values():
        if not q.children:
            _question_sheet(wb, q, units, questions_sheet)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return problems


def gather_units(settings) -> tuple[list[Unit], list[str]]:
    """Every included city's agencies: from its Sources workbook (the team's reviewed list,
    "Include?" = Yes) when Step 1 has run, otherwise from the city register."""
    from asics_agent.cities import load_register
    from asics_agent.sources_workbook import read_sources_workbook, sources_path, workspace_for

    cities, _ = load_register(settings.city_register)
    units, notes = [], []
    for city in cities.values():
        workspace = workspace_for(settings.outputs_dir, city.name)
        path = sources_path(workspace, city.name)
        agencies = city.parastatals
        if path.exists():
            found, _, _ = read_sources_workbook(path, workspace / "evidence")
            agencies = [p for p in found if p.include] or agencies
        else:
            notes.append(
                f"{city.name}: no Sources workbook yet (Step 1 not run), so only the "
                "parastatals listed in the city register are included."
            )
        if not agencies:
            notes.append(f"{city.name}: no parastatals known yet, so it has no rows.")
        units += [Unit(city.name, p) for p in agencies]
    return units, notes
