"""Human scoring: assign questions to interns, give each a copy, merge their work back.

Everything is Excel. The research lead fills in the assignments workbook:

  Verticals sheet     Vertical | Scoring workbook (a file in scoring/2027/templates)
  Assignments sheet   Vertical | Intern | Questions | Cities | Notes
                      Questions: codes separated by commas; "UPD1*" means every code that
                      starts with UPD1. Cities: names separated by commas, or "All".

One intern can score UPD1a for 10 cities and another the same UPD1a for 10 other cities.
Each intern gets their own copy of the vertical's workbook (scoring/2027/human/<vertical>/),
with their questions and rows highlighted and other question sheets hidden. Merging copies
only what interns type (inputs, evidence, comments), never formulas, into
scoring/2027/merged/<vertical>_human.xlsx, and lists anything that needs a look.
"""

import difflib
import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from asics_agent.scoring.template import QuestionSheet, VerticalTemplate, read_template

ASSIGNMENTS_FILE = "ASICS_2027_Scoring_Assignments.xlsx"
VERTICALS_SHEET, ASSIGNMENTS_SHEET = "Verticals", "Assignments"
VERTICAL_COLUMNS = ["Vertical", "Scoring workbook"]
ASSIGNMENT_COLUMNS = ["Vertical", "Intern", "Questions", "Cities", "Notes"]
APPLIES = "Applies?"
SEPARATOR = " – "  # "Bengaluru – BWSSB"
_HEAD = PatternFill("solid", fgColor="DDE7F0")
_MINE = PatternFill("solid", fgColor="FFF6D5")
_NOT_MINE = PatternFill("solid", fgColor="F2F2F2")


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _key(name: str) -> str:
    return re.sub(r"[^a-z0-9]", "", name.casefold())


@dataclass
class Assignment:
    vertical: str
    intern: str
    questions: list[str]  # resolved question codes
    rows: list[str]  # resolved row labels
    sheet_row: int


@dataclass
class Plan:
    verticals: dict[str, Path]  # vertical -> template path
    assignments: list[Assignment]
    problems: list[str] = field(default_factory=list)


def city_of(label: str) -> str:
    return label.split(SEPARATOR)[0].strip()


def match_cities(labels: list[str], cities) -> tuple[dict[str, str], list[str]]:
    """Map each row label's city to a register city (name or "Other names").

    Returns ({template city: register name}, problems)."""
    known = {}
    for city in cities.values():
        for name in [city.name, *city.aliases]:
            known[_key(name)] = city.name
    mapping, problems, unknown = {}, [], []
    for city in dict.fromkeys(city_of(label) for label in labels):
        if _key(city) in known:
            mapping[city] = known[_key(city)]
            continue
        guess = difflib.get_close_matches(city, [c.name for c in cities.values()], n=1, cutoff=0.6)
        if guess:
            problems.append(
                f'The scoring workbook\'s "{city}" looks like {guess[0]} in the city '
                f'register. If it is, add "{city}" to {guess[0]}\'s "Other names".'
            )
        else:
            unknown.append(city)
    if unknown:
        problems.append(
            f"{len(unknown)} cities in the scoring workbook aren't in the city "
            f"register yet ({', '.join(unknown)}). They can be scored by people; "
            "add them to the register before the AI scores them."
        )
    return mapping, problems


def write_assignments_template(path: Path, verticals: dict[str, str]) -> Path:
    """An empty assignments workbook with a guide, ready for the research lead."""
    wb = Workbook()
    guide = wb.active
    guide.title = "How to use"
    guide.column_dimensions["A"].width = 110
    for line in [
        "Scoring assignments",
        "",
        "Verticals sheet: one row per vertical and the scoring workbook it uses (a file in "
        "scoring/2027/templates).",
        "Assignments sheet: one row per intern and set of questions.",
        '  Questions: codes separated by commas, e.g. "UPD1a, UPD1b". "UPD1*" means every '
        "code starting with UPD1.",
        '  Cities: city names separated by commas, or "All". To split a question between '
        "interns, give each a different set of cities.",
        "After saving, open the app's Scoring page (or run scripts/scoring.py prepare) to make "
        "each intern's copy. Interns work only in their own copy.",
    ]:
        guide.append([line])
    guide["A1"].font = Font(bold=True, size=14)
    for title, columns, rows in [
        (VERTICALS_SHEET, VERTICAL_COLUMNS, list(verticals.items())),
        (ASSIGNMENTS_SHEET, ASSIGNMENT_COLUMNS, []),
    ]:
        ws = wb.create_sheet(title)
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), _HEAD
            ws.column_dimensions[cell.column_letter].width = 30
        for row in rows:
            ws.append(list(row))
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _sheet_rows(ws, columns: list[str], problems: list[str]):
    header = [_text(c.value) for c in ws[1]]
    missing = [c for c in columns if c not in header]
    if missing:
        problems.append(f"The {ws.title} sheet is missing the column(s) {', '.join(missing)}.")
        return []
    index = {c: header.index(c) for c in columns}
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = {c: _text(row[i]) if i < len(row) else "" for c, i in index.items()}
        if any(values.values()):
            yield r, values


def _code_matches(code: str, pattern: str) -> bool:
    """ "UPD 1a" matches UPD1a; "UPD1*" matches UPD1, UPD1a, UPD1b but not UPD10."""
    wanted = re.sub(r"[^a-z0-9*]", "", pattern.casefold())
    if not wanted.endswith("*"):
        return _key(code) == wanted
    prefix = wanted.rstrip("*")
    rest = _key(code)[len(prefix) :]
    return _key(code).startswith(prefix) and not (prefix[-1:].isdigit() and rest[:1].isdigit())


def _split(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"[,;\n]", text) if p.strip()]


def read_plan(path: Path, templates_dir: Path, register=None) -> tuple[Plan, dict]:
    """Read the assignments workbook; returns (plan, {vertical: VerticalTemplate})."""
    problems: list[str] = []
    wb = load_workbook(path, data_only=True)
    for sheet in (VERTICALS_SHEET, ASSIGNMENTS_SHEET):
        if sheet not in wb.sheetnames:
            return Plan({}, [], [f'The assignments workbook has no "{sheet}" sheet.']), {}
    verticals, templates = {}, {}
    for r, v in _sheet_rows(wb[VERTICALS_SHEET], VERTICAL_COLUMNS, problems):
        file = templates_dir / v["Scoring workbook"]
        if not file.exists():
            problems.append(
                f"Verticals sheet, row {r}: there is no file "
                f'"{v["Scoring workbook"]}" in {templates_dir}.'
            )
            continue
        verticals[v["Vertical"].upper()] = file
        templates[v["Vertical"].upper()] = read_template(file)

    assignments = []
    for r, v in _sheet_rows(wb[ASSIGNMENTS_SHEET], ASSIGNMENT_COLUMNS, problems):
        where = f"Assignments sheet, row {r}"
        vertical = v["Vertical"].upper()
        if vertical not in templates:
            problems.append(
                f'{where}: the vertical "{v["Vertical"]}" isn\'t on the Verticals sheet.'
            )
            continue
        if not v["Intern"]:
            problems.append(f"{where}: add the intern's name.")
            continue
        template = templates[vertical]
        codes = []
        for pattern in _split(v["Questions"]) or ["*"]:
            found = [c for c in template.questions if _code_matches(c, pattern)]
            if not found:
                problems.append(f'{where}: no question "{pattern}" in the {vertical} workbook.')
            codes += [c for c in found if c not in codes]
        rows = _rows_for(template, v["Cities"], register, where, problems)
        if codes and rows:
            assignments.append(Assignment(vertical, v["Intern"], codes, rows, r))
    problems += _overlaps(assignments)
    return Plan(verticals, assignments, problems), templates


def _rows_for(
    template: VerticalTemplate, cities_text: str, register, where: str, problems: list[str]
) -> list[str]:
    wanted = _split(cities_text)
    if not wanted or [w.lower() for w in wanted] == ["all"]:
        return list(template.labels)
    aliases = {}
    if register:
        for city in register.values():
            for name in [city.name, *city.aliases]:
                aliases[_key(name)] = _key(city.name)
    canon = lambda name: aliases.get(_key(name), _key(name))  # noqa: E731
    rows = []
    for city in wanted:
        found = [label for label in template.labels if canon(city_of(label)) == canon(city)]
        if not found:
            problems.append(f'{where}: the city "{city}" has no rows in the workbook.')
        rows += found
    return rows


def _overlaps(assignments: list[Assignment]) -> list[str]:
    seen, problems = {}, []
    for a in assignments:
        for code in a.questions:
            for row in a.rows:
                other = seen.setdefault((a.vertical, code, row), a.intern)
                if other != a.intern:
                    problems.append(
                        f"{code} for {row} is assigned to both {other} and "
                        f"{a.intern}. Their answers will be compared when merging."
                    )
    return problems[:50]


def intern_file(scoring_dir: Path, vertical: str, intern: str) -> Path:
    safe = re.sub(r"[^A-Za-z0-9]+", "_", intern).strip("_")
    return scoring_dir / "human" / vertical / f"{vertical}_{safe}.xlsx"


def prepare_intern_copies(
    plan: Plan, templates: dict, scoring_dir: Path, evidence_for=None
) -> list[str]:
    """Make each intern's workbook. Existing copies are never overwritten (they hold work).

    `evidence_for(vertical, code, label)` may return the Step 2 evidence for a row (a dict
    of column -> text); it is listed on a "Step 2 evidence" sheet in the copy."""
    notes = []
    by_intern: dict[tuple[str, str], list[Assignment]] = {}
    for a in plan.assignments:
        by_intern.setdefault((a.vertical, a.intern), []).append(a)
    for (vertical, intern), items in by_intern.items():
        path = intern_file(scoring_dir, vertical, intern)
        if path.exists():
            notes.append(f"{intern} already has a {vertical} copy; left as it is.")
            continue
        path.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(plan.verticals[vertical], path)
        rows_evidence = (
            (lambda code, label, v=vertical: evidence_for(v, code, label)) if evidence_for else None
        )
        _mark_copy(path, templates[vertical], intern, items, rows_evidence)
        notes.append(f"Made {path.name} for {intern}.")
    return notes


def _mark_copy(
    path: Path, template: VerticalTemplate, intern: str, items: list[Assignment], evidence_for=None
):
    wb = load_workbook(path)
    mine = {code: set() for a in items for code in a.questions}
    for a in items:
        for code in a.questions:
            mine[code].update(a.rows)
    for code, sheet in template.questions.items():
        ws = wb[sheet.sheet]
        if code not in mine:
            ws.sheet_state = "hidden"
            continue
        for label, r in sheet.rows.items():
            ws.cell(row=r, column=1).fill = _MINE if label in mine[code] else _NOT_MINE
    ws = wb.create_sheet("Your assignment", 0)
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 90
    ws.append([f"{intern}: your questions"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(
        [
            "Fill in only the rows highlighted in yellow on each question sheet. Enter what "
            "the row above each column asks for, the evidence (document, provision, page, "
            "link) and your comments. Don't change formulas or other rows."
        ]
    )
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.append([])
    ws.append(["Question", "Rows to score"])
    for cell in ws[4]:
        cell.font, cell.fill = Font(bold=True), _HEAD
    for code, rows in mine.items():
        ordered = [label for label in template.labels if label in rows]
        ws.append([code, ", ".join(ordered)])
        ws.cell(row=ws.max_row, column=1).hyperlink = f"#'{template.questions[code].sheet}'!A1"
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(wrap_text=True)
    ws.append([])
    ws.append(
        [
            "Note",
            "Scores that compare cities with each other are only final after "
            "everyone's work is merged.",
        ]
    )
    if evidence_for:
        _evidence_sheet(wb, template, mine, evidence_for)
    wb.active = 0
    wb.save(path)


def _evidence_sheet(wb, template: VerticalTemplate, mine: dict, evidence_for) -> None:
    """What Step 2 found for each assigned row: score from these citations."""
    rows = []
    for code, labels in mine.items():
        for label in template.labels:
            if label in labels and (found := evidence_for(code, label)):
                rows.append({"Question": code, "Row": label, **found})
    if not rows:
        return
    ws = wb.create_sheet("Step 2 evidence", 1)
    ws.append(
        [
            "Score from these Step 2 citations (the reviewed Citation Sheet). If one doesn't "
            "support the answer, say so in SCORER COMMENTS."
        ]
    )
    ws["A1"].font = Font(bold=True)
    columns = list(rows[0])
    ws.append(columns)
    for cell in ws[2]:
        cell.font, cell.fill = Font(bold=True), _HEAD
        ws.column_dimensions[cell.column_letter].width = 22
    for row in rows:
        ws.append([row.get(c, "") for c in columns])
        link = row.get("Link")
        if link:
            ws.cell(row=ws.max_row, column=columns.index("Link") + 1).hyperlink = link
    for name in ("Quote", "Step 2 answer"):
        if name in columns:
            ws.column_dimensions[
                ws.cell(row=2, column=columns.index(name) + 1).column_letter
            ].width = 70
    ws.freeze_panes = "C3"


def _typed_cells(sheet: QuestionSheet) -> list[int]:
    """Columns an intern types into: evidence, inputs, comments (never the Points formula)."""
    cols = [*sheet.evidence.values(), *(i.column for i in sheet.inputs)]
    if sheet.comments_column:
        cols.append(sheet.comments_column)
    return sorted(set(cols) - {sheet.points_column})


def _input_problem(sheet: QuestionSheet, column: int, value) -> str | None:
    spec = next((i for i in sheet.inputs if i.column == column), None)
    if spec is None or value in (None, ""):
        return None
    text = _text(value)
    if spec.kind == "choice":
        options = [o.upper() for o in spec.options]
        try:
            ok = text.upper() in options or any(float(text) == float(o) for o in options)
        except ValueError:
            ok = False
        return None if ok else f"should be one of {' / '.join(spec.options)}"
    if spec.kind in ("whole_number", "number", "percent"):
        try:
            number = float(text.rstrip("%"))
        except ValueError:
            return "should be a number"
        if spec.kind == "whole_number" and number != int(number):
            return "should be a whole number"
    return None


@dataclass
class MergeResult:
    path: Path
    copied: int = 0  # cells copied from interns
    done: dict[str, tuple[int, int]] = field(default_factory=dict)  # intern -> (done, total)
    problems: list[str] = field(default_factory=list)


def _applies(ws, sheet: QuestionSheet, row: int) -> bool:
    for c in range(1, ws.max_column + 1):
        if _text(ws.cell(row=sheet.header_row, column=c).value) == APPLIES:
            return _text(ws.cell(row=row, column=c).value).upper() != "NO"
    return True


def merge_vertical(plan: Plan, templates: dict, vertical: str, scoring_dir: Path) -> MergeResult:
    """Copy every intern's typed cells for their assigned rows into a fresh copy of the
    template. Conflicting answers (two interns, same cell, different values) are listed and
    the first intern's value is kept."""
    template: VerticalTemplate = templates[vertical]
    out = scoring_dir / "merged" / f"{vertical}_human.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(plan.verticals[vertical], out)
    merged = load_workbook(out)
    result = MergeResult(out)
    written: dict[tuple[str, int, int], tuple[str, object]] = {}
    for a in (a for a in plan.assignments if a.vertical == vertical):
        path = intern_file(scoring_dir, vertical, a.intern)
        if not path.exists():
            result.problems.append(f"{a.intern} has no {vertical} copy yet.")
            continue
        source = load_workbook(path)
        done, total = result.done.get(a.intern, (0, 0))
        for code in a.questions:
            sheet = template.questions[code]
            src, dst = source[sheet.sheet], merged[sheet.sheet]
            for label in a.rows:
                r = sheet.rows[label]
                if not _applies(dst, sheet, r):
                    continue
                total += 1
                inputs = [_text(src.cell(row=r, column=i.column).value) for i in sheet.inputs]
                done += all(inputs)
                for c in _typed_cells(sheet):
                    value = src.cell(row=r, column=c).value
                    if value in (None, "") or _text(value).startswith("="):
                        continue
                    if problem := _input_problem(sheet, c, value):
                        result.problems.append(f'{a.intern}, {code}, {label}: "{value}" {problem}.')
                    key = (sheet.sheet, r, c)
                    if key in written and _text(written[key][1]) != _text(value):
                        other, kept = written[key]
                        result.problems.append(
                            f'{code}, {label}: {other} entered "{kept}" and {a.intern} '
                            f'entered "{value}". Kept {other}\'s; please agree on one.'
                        )
                        continue
                    written.setdefault(key, (a.intern, value))
                    dst.cell(row=r, column=c, value=value)
                    result.copied += 1
        result.done[a.intern] = (done, total)
    result.problems = list(dict.fromkeys(result.problems))
    _report_sheet(merged, result, template)
    merged.save(out)
    return result


def _report_sheet(wb, result: MergeResult, template: VerticalTemplate):
    ws = wb.create_sheet("Merge report", 0)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100
    ws.append(["Merged from the interns' copies"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Intern", "Rows done / assigned"])
    for intern, (done, total) in result.done.items():
        ws.append([intern, f"{done} / {total}"])
    ws.append([])
    ws.append(["Needs a look", f"{len(result.problems)} item(s)"])
    for p in result.problems:
        ws.append(["", p])
    for row in ws.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True)
    wb.active = 0
