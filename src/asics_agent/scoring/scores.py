"""Reads the scores out of a calculated scoring workbook and builds the results workbook.

Every vertical's workbook has the same summary layout (see template.py), so the same code
reads people's scores (merged copy) and the AI's (AI copy) for any vertical:
question/sub-question scores per row, and the overall score per row. Rows are cities, or
"City – Agency"; a city's score is the average of its rows (equal weights).

The results workbook (scoring/2027/results/ASICS_2027_Scores.xlsx) shows, per city, each
vertical's score from people and from the AI, and the ASICS score as an Excel formula
(the equal-weight average of the verticals), plus every question score side by side.
"""

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from asics_agent.scoring.human import city_of
from asics_agent.scoring.template import VerticalTemplate

_HEAD = PatternFill("solid", fgColor="DDE7F0")
_PROVISIONAL = PatternFill("solid", fgColor="FFF6D5")
Value = float | str | None  # a number, "NA", "?" (not scored yet), "ER" (sheet problem)


def _value(v) -> Value:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).strip()
    try:
        return float(text)
    except ValueError:
        return text


@dataclass
class VerticalScores:
    vertical: str
    source: str  # "People" or "AI"
    path: Path
    overall: dict[str, Value] = field(default_factory=dict)  # row label -> overall score
    not_scored: dict[str, int] = field(default_factory=dict)  # row label -> MQs still "?"
    questions: dict[tuple[str, str], Value] = field(default_factory=dict)  # (code, label)
    tags: dict[str, str] = field(default_factory=dict)  # code -> MQ / SQ

    def city_scores(self, canon=None) -> dict[str, tuple[Value, bool]]:
        """{city: (score, provisional)}: the average of the city's rows that have a score.
        `canon` maps a workbook's city name to the register's (e.g. Bangalore -> Bengaluru)."""
        cities: dict[str, list[str]] = {}
        for label in self.overall:
            city = city_of(label)
            cities.setdefault(canon(city) if canon else city, []).append(label)
        out = {}
        for city, labels in cities.items():
            numbers = [self.overall[x] for x in labels if isinstance(self.overall[x], float)]
            provisional = any(self.not_scored.get(x) for x in labels)
            if numbers:
                out[city] = (sum(numbers) / len(numbers), provisional)
            else:
                values = {self.overall[x] for x in labels}
                out[city] = ("NA" if values <= {"NA"} else "?", provisional)
        return out


def read_scores(
    template: VerticalTemplate, path: Path, vertical: str, source: str
) -> VerticalScores:
    """Scores from a workbook that has been calculated (see recalc.py)."""
    summary = template.summary
    ws = load_workbook(path, data_only=True)[summary.sheet]
    scores = VerticalScores(vertical, source, path)
    tag_col = next(
        (
            c
            for c in range(1, 10)
            if str(ws.cell(row=summary.header_row, column=c).value or "").strip() == "MQ_SQ"
        ),
        None,
    )
    for code, row in summary.question_rows.items():
        if tag_col:
            scores.tags[code] = str(ws.cell(row=row, column=tag_col).value or "")
        for label, col in summary.labels.items():
            scores.questions[(code, label)] = _value(ws.cell(row=row, column=col).value)
    if summary.overall_row:
        for label, col in summary.labels.items():
            scores.overall[label] = _value(ws.cell(row=summary.overall_row, column=col).value)
            count = _value(ws.cell(row=summary.overall_row + 1, column=col).value)
            scores.not_scored[label] = int(count) if isinstance(count, float) else 0
    mains = [code for code, tag in scores.tags.items() if tag == "MQ"]
    for label in scores.overall:  # nothing scored yet: "?", not the formula's 0
        values = [scores.questions.get((code, label)) for code in mains]
        if mains and not any(isinstance(v, float) for v in values):
            scores.overall[label] = "NA" if all(v == "NA" for v in values) else "?"
    return scores


def _round(v: Value):
    return round(v, 2) if isinstance(v, float) else v


def build_results(scores: list[VerticalScores], path: Path, canon=None) -> Path:
    verticals = list(dict.fromkeys(s.vertical for s in scores))
    by = {(s.vertical, s.source): s for s in scores}
    wb = Workbook()
    ws = wb.active
    ws.title = "City scores"
    columns = []  # (vertical, source)
    for v in verticals:
        columns += [(v, src) for src in ("People", "AI") if (v, src) in by]
    header = ["City", *[f"{v} ({src})" for v, src in columns]]
    sources = [src for src in ("People", "AI") if any(c[1] == src for c in columns)]
    header += [f"ASICS score ({src})" for src in sources]
    header += [f"Main questions not yet scored ({src})" for src in sources]
    ws.append(header)
    city_values = {c: by[c].city_scores(canon) for c in columns}
    cities = list(dict.fromkeys(city for c in columns for city in city_values[c]))
    for r, city in enumerate(cities, start=2):
        ws.cell(row=r, column=1, value=city)
        for i, c in enumerate(columns, start=2):
            value, provisional = city_values[c].get(city, (None, False))
            cell = ws.cell(row=r, column=i, value=_round(value))
            if provisional:
                cell.fill = _PROVISIONAL
        for k, src in enumerate(sources):
            cells = [
                f"{get_column_letter(i)}{r}" for i, c in enumerate(columns, start=2) if c[1] == src
            ]
            ws.cell(
                row=r,
                column=len(columns) + 2 + k,
                value=f'=IFERROR(ROUND(AVERAGE({",".join(cells)}),2),"")',
            )
            pending = sum(
                n
                for c in columns
                if c[1] == src
                for label, n in by[c].not_scored.items()
                if (canon(city_of(label)) if canon else city_of(label)) == city
            )
            ws.cell(row=r, column=len(columns) + 2 + len(sources) + k, value=pending)
    note = len(cities) + 3
    ws.cell(
        row=note,
        column=1,
        value=(
            "ASICS score = the average of the city's vertical scores (equal weights). Yellow = "
            "provisional: some main questions aren't scored yet. NA = no question applies."
        ),
    )
    ws.cell(row=note, column=1).font = Font(italic=True)

    long = wb.create_sheet("Question scores")
    long.append(
        [
            "Vertical",
            "Question",
            "MQ_SQ",
            "Row",
            "City",
            "People",
            "AI",
            "Difference (AI - People)",
            "Same score?",
        ]
    )
    for v in verticals:
        people, ai = by.get((v, "People")), by.get((v, "AI"))
        base = people or ai
        for code, label in base.questions:
            p = people.questions.get((code, label)) if people else None
            a = ai.questions.get((code, label)) if ai else None
            both = isinstance(p, float) and isinstance(a, float)
            long.append(
                [
                    v,
                    code,
                    base.tags.get(code, ""),
                    label,
                    city_of(label),
                    _round(p),
                    _round(a),
                    round(a - p, 2) if both else None,
                    ("Yes" if abs(a - p) < 1e-9 else "No")
                    if both
                    else ("Yes" if p == a and p in ("NA",) else None),
                ]
            )
    for sheet in (ws, long):
        for cell in sheet[1]:
            cell.font, cell.fill = Font(bold=True), _HEAD
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            sheet.column_dimensions[cell.column_letter].width = 16
        sheet.column_dimensions["A"].width = 20
        sheet.freeze_panes = "B2"
    long.auto_filter.ref = long.dimensions
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def agreement(people: VerticalScores, ai: VerticalScores, leaf_codes: set[str]) -> dict:
    """How closely the AI matches people (the golden dataset) on the questions people score
    directly (not the roll-ups)."""
    pairs = [
        (people.questions[k], ai.questions.get(k)) for k in people.questions if k[0] in leaf_codes
    ]
    both = [(p, a) for p, a in pairs if isinstance(p, float) and isinstance(a, float)]
    people_only = sum(1 for p, a in pairs if isinstance(p, float) and not isinstance(a, float))
    same = sum(1 for p, a in both if abs(p - a) < 1e-9)
    return {
        "Vertical": people.vertical,
        "Scored by both": len(both),
        "Same score": same,
        "Agreement %": round(100 * same / len(both), 1) if both else None,
        "Average difference": round(sum(abs(p - a) for p, a in both) / len(both), 2)
        if both
        else None,
        "Scored by people only": people_only,
    }


def add_eval_history(rows: list[dict], path: Path, label: str) -> Path:
    """Append this comparison to the evals workbook, so changes to the AI can be compared."""
    from datetime import datetime

    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "AI vs people"
        ws.append(["When", "Label", *rows[0].keys()] if rows else ["When", "Label"])
        for cell in ws[1]:
            cell.font, cell.fill = Font(bold=True), _HEAD
            ws.column_dimensions[cell.column_letter].width = 18
    when = datetime.now().strftime("%Y-%m-%d %H:%M")
    for row in rows:
        ws.append([when, label, *row.values()])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
