"""What Step 3 (scoring) needs from Steps 1 and 2: is a city ready, and its evidence.

A city is ready for Step 3 when Step 2 has answered it after the last change to its Sources
workbook. Its evidence is the latest answers workbook (one answer per agency and question,
naming the Citation ID it relied on) and the reviewed Citation Sheet in the Sources workbook.
"""

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from asics_agent.models import Parastatal, Source
from asics_agent.sources_workbook import read_sources_workbook, sources_path, workspace_for
from asics_agent.workbook import answers_pattern, is_scoring_sheet

NO_EVIDENCE = {"Insufficient Evidence", "Not Yet Assessed", ""}


@dataclass
class Phase2:
    city: str
    ready: bool
    reason: str  # plain language: what is missing, or when Step 2 answered it
    sources: Path | None = None
    answers: Path | None = None
    answered_at: datetime | None = None


@dataclass
class StepTwoAnswer:
    status: str
    answer: str
    citation_id: str
    evidence: str
    notes: str
    source_needed: str
    question: str = ""  # the question's text as Step 2 answered it


@dataclass
class CityEvidence:
    parastatals: dict[str, Parastatal] = field(default_factory=dict)  # id -> parastatal
    citations: list[Source] = field(default_factory=list)  # usable Citation Sheet rows
    answers: dict[tuple[str, str], StepTwoAnswer] = field(default_factory=dict)  # (id, code)


def phase2_status(settings, city: str) -> Phase2:
    """Is `city` ready for Step 3 in this vertical (settings: the vertical's)?"""
    workspace = workspace_for(settings.outputs_dir, city)
    sources = sources_path(workspace, city)
    if not sources.exists():
        return Phase2(city, False, "Step 1 (find sources) hasn't been run for this city.")
    answers = sorted(
        (workspace / "answers").glob(f"*/{answers_pattern(settings)}"),
        key=lambda p: p.stat().st_mtime,
    )
    if not answers:
        return Phase2(
            city, False, "Step 2 (answer questions) hasn't been run for this city.", sources
        )
    latest = answers[-1]
    when = datetime.fromtimestamp(latest.stat().st_mtime)
    if latest.stat().st_mtime < sources.stat().st_mtime:
        return Phase2(
            city,
            False,
            "The Sources workbook changed after Step 2. Run Step 2 "
            "again so the answers use the current Citation Sheet.",
            sources,
            latest,
            when,
        )
    return Phase2(city, True, f"Answered {when:%d %b %Y %H:%M}.", sources, latest, when)


def usable(source: Source) -> bool:
    """The same rule Step 2 uses: reviewed, verified, and its saved text is on disk."""
    return (
        source.use_for_answers
        and source.verification_status == "Verified"
        and bool(source.content_path)
        and Path(source.content_path).exists()
    )


def _code(question_id: str) -> str:
    """ "UPD 1a" / "SC-1" -> "UPD1a" / "SC1": the scoring workbook's sheet codes."""
    return re.sub(r"[\s\-]", "", question_id)


def bank_ids(settings) -> dict[str, str]:
    """{the bank's own ID: normalised ID} for the vertical's question bank."""
    from asics_agent.verticals import load_bank

    questions, _ = load_bank(settings)
    return {q.raw_id: q.id for q in questions.values()}


def load_evidence(status: Phase2, ids: dict[str, str] | None = None) -> CityEvidence:
    """The city's usable citations and Step 2 answers, keyed by (unit ID, question code).

    `ids` maps the question bank's own IDs (e.g. "SC-1", "DPG d") to normalised ones
    ("SC 1", "DPG 5d"); the answers workbook keeps the bank's own spelling."""
    ids = ids or {}
    workspace = status.sources.parent
    parastatals, citations, _ = read_sources_workbook(status.sources, workspace / "evidence")
    evidence = CityEvidence(
        parastatals={p.id: p for p in parastatals if p.include},
        citations=[c for c in citations if usable(c)],
    )
    by_name = {p.name.casefold(): p.id for p in parastatals}
    by_name.update({p.id.casefold(): p.id for p in parastatals})
    wb = load_workbook(status.answers, read_only=True, data_only=True)
    sheet = next((ws for ws in wb.worksheets if is_scoring_sheet(ws.title)), None)
    if sheet is None:
        return evidence
    rows = sheet.iter_rows(values_only=True)
    header = [str(h or "").strip() for h in next(rows)]
    col = {h: i for i, h in enumerate(header)}

    def get(row, name):
        i = col.get(name)
        return "" if i is None or i >= len(row) or row[i] is None else str(row[i]).strip()

    for row in rows:
        pid = by_name.get(get(row, "Parastatal").casefold())
        qid = get(row, "Question ID")
        if pid and qid:
            evidence.answers[(pid, _code(ids.get(qid, qid)))] = StepTwoAnswer(
                question=get(row, "Question"),
                status=get(row, "Status"),
                answer=get(row, "Answer"),
                citation_id=get(row, "Citation ID"),
                evidence=get(row, "Evidence Excerpt"),
                notes=get(row, "Notes"),
                source_needed=get(row, "Source Needed"),
            )
    return evidence
