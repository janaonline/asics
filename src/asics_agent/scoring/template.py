"""Reads a vertical's scoring workbook (the expert team's template) by its conventions.

The scoring criteria live in the workbook itself (the experts' formulas); this module only
finds where things are, so AI and humans can fill the same cells and the tool can read the
results. Conventions, as in the ASICS 2027 UPD workbook:

  Question sheets   A1 = "Question Code", B1 = the code (= sheet name). Rows above the header
                    show the question details; the row just above the header says what to
                    enter in each column ("SELECT YES / NO", "ENTER A WHOLE NUMBER",
                    "DO NOT ALTER FORMULA", "AUTO - DO NOT ENTER"…). The header row starts
                    with "City"; one row per city (or city – agency) below it.
  Columns           evidence (Act Name/Web name/Doc name, Chapter, Provision, Clause, Page
                    Number, Link), inputs, "Points (Auto-generate)", "SCORER COMMENTS".
  Summary sheet     header row with the question-code column and one column per city; one
                    row per question code; an overall-score row (label ending "_OVERALL_SCORE").
"""

import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

POINTS = "Points (Auto-generate)"
COMMENTS = "SCORER COMMENTS"
EVIDENCE = ["Act Name/Web name/Doc name", "Chapter", "Provision", "Clause", "Page Number", "Link"]
_OPTIONS = re.compile(r"^SELECT\s+(.+)$", re.I)


@dataclass
class InputColumn:
    column: int
    header: str
    instruction: str  # the text above the header, e.g. "SELECT YES / NO"
    kind: str  # "choice" | "whole_number" | "number" | "percent" | "text"
    options: list[str] = field(default_factory=list)  # for choices, e.g. ["YES", "NO"]


@dataclass
class QuestionSheet:
    code: str
    sheet: str
    header_row: int
    rows: dict[str, int]  # row label (city, or "City – Agency") -> row number
    evidence: dict[str, int]  # evidence header -> column
    inputs: list[InputColumn]
    points_column: int | None
    comments_column: int | None
    details: dict[str, int]  # question detail label (A column above the header) -> row


@dataclass
class SummarySheet:
    sheet: str
    header_row: int
    code_column: int
    labels: dict[str, int]  # city / unit label -> column
    question_rows: dict[str, int]  # question code -> row
    overall_row: int | None
    overall_label: str = ""


@dataclass
class VerticalTemplate:
    path: Path
    questions: dict[str, QuestionSheet]
    summary: SummarySheet | None
    questions_list_sheet: str | None
    labels: list[str]  # the row labels (cities, or city – agency), in template order
    problems: list[str] = field(default_factory=list)  # plain-language template problems


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def readable(header: str) -> str:
    """Header text a person sees, from a header that may be a formula like
    ="Part A: Board provided ("&TEXT($C$7/2,"0.00")&")"  ->  Part A: Board provided"""
    if not header.startswith('="'):
        return header
    text = header[2:].split('"&', 1)[0].rstrip('"')
    return re.sub(r"\s*\($", "", text).strip()


def _kind(instruction: str) -> tuple[str, list[str]]:
    upper = instruction.upper()
    if match := _OPTIONS.match(instruction):
        return "choice", [o.strip().upper() for o in match.group(1).split("/") if o.strip()]
    if re.fullmatch(r"[A-Z0-9 ]+(/ *[A-Z0-9 ]+)+", upper):  # e.g. "SHALL / MAY / NONE"
        return "choice", [o.strip() for o in upper.split("/")]
    if "WHOLE NUMBER" in upper or "POPULATION" in upper:
        return "whole_number", []
    if "%" in upper:
        return "percent", []
    if "NUMBER" in upper or "AMOUNT" in upper:
        return "number", []
    return "text", []


def _is_input(instruction: str) -> bool:
    upper = instruction.upper()
    return bool(upper) and not any(k in upper for k in ("DO NOT", "AUTO"))


def read_question_sheet(ws: Worksheet) -> QuestionSheet | None:
    if _text(ws["A1"].value) != "Question Code":
        return None
    header = next(
        (r for r in range(2, 40) if _text(ws.cell(row=r, column=1).value) == "City"), None
    )
    if header is None:
        return None
    headers = {c: _text(ws.cell(row=header, column=c).value) for c in range(1, ws.max_column + 1)}
    marks = {c: _text(ws.cell(row=header - 1, column=c).value) for c in headers}
    inputs = []
    for c, instruction in marks.items():
        if _is_input(instruction):
            kind, options = _kind(instruction)
            inputs.append(
                InputColumn(c, readable(headers[c]) or f"Column {c}", instruction, kind, options)
            )
    rows = {}
    for r in range(header + 1, ws.max_row + 1):
        label = _text(ws.cell(row=r, column=1).value)
        if not label:
            break
        rows[label] = r
    return QuestionSheet(
        code=_text(ws["B1"].value) or ws.title,
        sheet=ws.title,
        header_row=header,
        rows=rows,
        evidence={h: c for c, h in headers.items() if h in EVIDENCE},
        inputs=inputs,
        points_column=next((c for c, h in headers.items() if h == POINTS), None),
        comments_column=next((c for c, h in headers.items() if h.upper() == COMMENTS), None),
        details={
            _text(ws.cell(row=r, column=1).value): r
            for r in range(1, header - 1)
            if _text(ws.cell(row=r, column=1).value)
        },
    )


def _read_summary(wb) -> SummarySheet | None:
    for ws in wb.worksheets:
        if "SUMMARY" not in ws.title.upper():
            continue
        for r in range(1, 10):
            first = _text(ws.cell(row=r, column=1).value)
            if "Q No" in first or first.lower().startswith("question"):
                labels, codes, overall, overall_label = {}, {}, None, ""
                started = False
                for c in range(2, ws.max_column + 1):
                    head = _text(ws.cell(row=r, column=c).value)
                    if head == "Max Score":
                        started = True
                        continue
                    if started and head:
                        labels[head] = c
                for row in range(r + 1, ws.max_row + 1):
                    code = _text(ws.cell(row=row, column=1).value)
                    if code and code.upper() not in {"ENTER"}:
                        codes[code] = row
                    for c in range(1, 10):
                        text = _text(ws.cell(row=row, column=c).value)
                        if text.upper().endswith("_OVERALL_SCORE"):
                            overall, overall_label = row, text
                return SummarySheet(ws.title, r, 1, labels, codes, overall, overall_label)
    return None


def read_template(path: Path) -> VerticalTemplate:
    wb = load_workbook(path)  # formulas, not values: we need the structure
    questions, problems = {}, []
    for ws in wb.worksheets:
        sheet = read_question_sheet(ws)
        if sheet is None:
            continue
        if sheet.code != ws.title:
            problems.append(
                f"Sheet '{ws.title}' says its question code is '{sheet.code}'. The "
                "sheet name and the code in B1 must match."
            )
        if sheet.points_column is None:
            problems.append(f"Sheet '{ws.title}' has no '{POINTS}' column.")
        if not sheet.inputs:
            problems.append(
                f"Sheet '{ws.title}' has no input columns (the row above the "
                "headers should say what to enter, e.g. 'SELECT YES / NO')."
            )
        questions[sheet.code] = sheet
    summary = _read_summary(wb)
    if summary is None:
        problems.append("No summary sheet found (a sheet with SUMMARY in its name).")
    else:
        ws = wb[summary.sheet]
        first_label_col = min(summary.labels.values(), default=None)
        for code, row in summary.question_rows.items():
            formula = (
                _text(ws.cell(row=row, column=first_label_col).value) if first_label_col else ""
            )
            # e.g. INDIRECT("'"&$A52&"'!$B$1"): reads the sheet named by this row's code
            # (rows like UPD23a1 read parts of sheet UPD23a with LEFT(), which is fine)
            reads_own_sheet = f'"\'"&$A{row}&"\'!' in formula
            if reads_own_sheet and code not in questions:
                problems.append(
                    f"The summary reads question {code} from a sheet called "
                    f"'{code}', but there is no such sheet."
                )
    if summary is not None and summary.overall_row is None:
        problems.append(
            f"The summary sheet '{summary.sheet}' has no overall score row (a "
            "label ending in _OVERALL_SCORE)."
        )
    list_sheet = next(
        (ws.title for ws in wb.worksheets if "QUESTIONS" in ws.title.upper().replace(" ", "")), None
    )
    first = next(iter(questions.values()), None)
    labels = list(first.rows) if first else []
    for sheet in questions.values():
        if list(sheet.rows) != labels:
            problems.append(
                f"Sheet '{sheet.sheet}' lists different cities/rows from "
                f"'{first.sheet}'. Every question sheet must have the same rows."
            )
    return VerticalTemplate(path, questions, summary, list_sheet, labels, problems)
