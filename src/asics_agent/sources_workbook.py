"""The Sources workbook: one per city, the controlled list of parastatals and citations.

Step 1 (Find sources) writes it; the research team reviews and adds to it in Excel; Step 2
(Answer questions) answers ONLY from the rows it marks "Use for answers? = Yes".

Re-running Step 1 keeps the team's work: their Include / Use-for-answers decisions, their
own rows, and the "Your Notes" columns. The previous version is copied to history/ first.
"""

import hashlib
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from asics_agent.cities import FLAGS, TYPE_LABELS, TYPES, YES
from asics_agent.models import ALL_PARASTATALS, Issue, Parastatal, Question, Source
from asics_agent.workbook import (
    _FILLS,
    _HEADER_FILL,
    Link,
    _add_dropdown,
    _new_sheet,
    _norm,
    _saved_copy,
    _write_issue_sheets,
    _write_rows,
)

PARASTATALS_SHEET, CITATIONS_SHEET, COVERAGE_SHEET = "Parastatals", "Citation Sheet", "Coverage"

PARASTATAL_COLUMNS = [
    ("Parastatal ID", []),
    ("Parastatal Name", []),
    ("Type", []),
    ("Include?", []),
    ("Found By", []),
    ("Why Included", []),
    ("Official Website", []),
    ("Website Check", []),
    ("Other Websites Checked", []),
    ("Governing Act", []),
    ("Current Status", []),
    *[(label, []) for label in FLAGS],
    ("Notes", []),
    ("Your Notes", []),
]
CITATION_COLUMNS = [
    ("Citation ID", []),
    ("Parastatal", []),
    ("Source Title", []),
    ("Source Type", []),
    ("Official URL", []),
    ("Saved Copy", []),
    ("What This Source Is Useful For", []),
    ("Question IDs", []),
    ("Use for Answers?", []),
    ("Verification Status", []),
    ("Accessibility", []),
    ("Current/Active Status", []),
    ("Authority Score", []),
    ("Found By", []),
    ("Notes", []),
    ("Your Notes", []),
    ("Last Checked", []),
]
COVERAGE_COLUMNS = [
    ("Parastatal", []),
    ("Question ID", []),
    ("Question", []),
    ("Citations", []),
    ("Coverage", []),
]

GUIDE = [
    ("Sources for {city}", ""),
    ("", ""),
    (
        "What this is",
        "The list of parastatals and sources for {city}. Step 2 (Answer "
        'questions) uses ONLY the Citation Sheet rows marked "Use for Answers? = Yes", so this '
        "workbook controls exactly what the answers can rely on.",
    ),
    (
        "1. Parastatals sheet",
        "Check each parastatal the agent found. Set Include? to No to "
        "leave one out. To add one it missed, add a row with its ID, name and Type; leave Found "
        "By empty. The Website Check column explains how much the official website can be "
        "trusted, point by point.",
    ),
    (
        "2. Citation Sheet",
        "Open each link (or its saved copy). Set Use for Answers? to No for "
        "anything unsuitable. To add a source you know, add a row: Parastatal (an ID from the "
        'Parastatals sheet, or "All parastatals"), Source Title, Official URL and What This '
        "Source Is Useful For. Leave Citation ID empty; it is filled in next time. Your links are "
        "opened and checked the same way as the agent's.",
    ),
    (
        "3. Coverage sheet",
        "Shows which questions have at least one source. Questions marked "
        '"No source yet" will be answered "Insufficient Evidence" unless you add a source.',
    ),
    ("Your Notes", "Write anything in the Your Notes columns; they are never overwritten."),
    ("When you're done", "Save and close this file, then run Step 2 (Answer questions)."),
]


def workspace_for(outputs_dir: Path, city_name: str) -> Path:
    return outputs_dir / "cities" / re.sub(r"[^\w\- ]", "", city_name).strip()


def sources_path(workspace: Path, city_name: str) -> Path:
    return workspace / f"ASICS_{city_name.replace(' ', '_')}_Sources.xlsx"


def evidence_paths(cache_dir: Path, url: str) -> tuple[str | None, str | None]:
    """The saved text and copy of a source (see links/accessibility.py naming)."""
    stem = hashlib.sha1(url.encode()).hexdigest()[:16]
    text = cache_dir / f"{stem}.txt"
    copy = next(
        (
            cache_dir / f"{stem}{ext}"
            for ext in (".pdf", ".html")
            if (cache_dir / f"{stem}{ext}").exists()
        ),
        None,
    )
    return (str(text) if text.exists() else None), (str(copy) if copy else None)


def assign_citation_ids(sources: list[Source]) -> None:
    """Give every source without one a stable ID like "BWSSB-07"."""
    used: dict[str, int] = {}
    for s in sources:
        if match := re.fullmatch(r"(.+)-(\d+)", s.citation_id or ""):
            used[match.group(1)] = max(used.get(match.group(1), 0), int(match.group(2)))
    for s in sources:
        if not s.citation_id:
            prefix = "ALL" if s.parastatal_id == ALL_PARASTATALS else s.parastatal_id
            used[prefix] = used.get(prefix, 0) + 1
            s.citation_id = f"{prefix}-{used[prefix]:02d}"


def _yes(value) -> bool:
    return str(value or "").strip().lower() in YES


def _cell_url(cell) -> str:
    if cell.hyperlink and str(cell.hyperlink.target or "").startswith("http"):
        return cell.hyperlink.target
    value = str(cell.value or "").strip()
    return value if value.startswith("http") else ""


def _rows(ws):
    header = {_norm(c.value): i for i, c in enumerate(ws[1]) if c.value}
    for row_no, row in enumerate(ws.iter_rows(min_row=2), start=2):

        def get(name, _row=row):
            i = header.get(_norm(name))
            return "" if i is None or _row[i].value is None else str(_row[i].value).strip()

        def url(name, _row=row):
            i = header.get(_norm(name))
            return "" if i is None else _cell_url(_row[i])

        if any(c.value not in (None, "") for c in row):
            yield row_no, get, url


def read_sources_workbook(path: Path, cache_dir: Path):
    """Return (parastatals, citations, issues) as the research team left them."""
    issues: list[Issue] = []

    def team(severity, message, ref=None):
        issues.append(
            Issue(
                phase="initial_checks", severity=severity, message=message, ref=ref, audience="team"
            )
        )

    wb = load_workbook(path, data_only=True)
    for sheet in (PARASTATALS_SHEET, CITATIONS_SHEET):
        if sheet not in wb.sheetnames:
            team(
                "error",
                f'The Sources workbook has no "{sheet}" sheet. Run Step 1 again to recreate it.',
            )
            return [], [], issues

    parastatals: list[Parastatal] = []
    for row_no, get, url in _rows(wb[PARASTATALS_SHEET]):
        pid = re.sub(r"\s+", "", get("Parastatal ID")).upper()
        where = f"Parastatals sheet, row {row_no}"
        if not pid or not get("Parastatal Name"):
            team("error", f"{where}: every parastatal needs an ID and a name.")
            continue
        kind = next(
            (code for label, code in TYPES.items() if label.lower() == get("Type").lower()), None
        )
        if kind is None:
            team("error", f"{where}: the Type of {pid} must be one of: {', '.join(TYPES)}.", pid)
            continue
        if any(p.id == pid for p in parastatals):
            team("error", f"{where}: the ID {pid} is used twice.", pid)
            continue
        flags = {
            field: (_yes(get(label)) if get(label) else default)
            for label, (field, default) in FLAGS.items()
        }
        parastatals.append(
            Parastatal(
                id=pid,
                name=get("Parastatal Name"),
                type=kind,
                include=not get("Include?") or _yes(get("Include?")),
                found_by="Agent" if get("Found By") == "Agent" else "Your team",
                why_included=get("Why Included"),
                official_website=url("Official Website"),
                governing_act=get("Governing Act"),
                current_status=get("Current Status"),
                notes=get("Notes"),
                team_notes=get("Your Notes"),
                **flags,
            )
        )

    ids = {p.id for p in parastatals}
    citations: list[Source] = []
    for row_no, get, url in _rows(wb[CITATIONS_SHEET]):
        where = f"Citation Sheet, row {row_no}"
        link = url("Official URL")
        owner = get("Parastatal")
        owner = ALL_PARASTATALS if owner.lower() == ALL_PARASTATALS.lower() else owner.upper()
        if not link:
            team("warning", f"{where}: has no web link starting with http, so it was skipped.")
            continue
        if owner != ALL_PARASTATALS and owner not in ids:
            team(
                "warning",
                f'{where}: "{get("Parastatal")}" isn\'t on the Parastatals sheet. '
                f'Use an ID from that sheet or "{ALL_PARASTATALS}".',
            )
            continue
        team_row = get("Found By") != "Agent"
        text_path, copy_path = evidence_paths(cache_dir, link)
        score = get("Authority Score")
        citations.append(
            Source(
                citation_id=get("Citation ID"),
                parastatal_id=owner,
                title=get("Source Title") or link,
                url=link,
                source_type=get("Source Type"),
                useful_for=get("What This Source Is Useful For"),
                question_ids=[q.strip() for q in get("Question IDs").split(",") if q.strip()],
                use_for_answers=not get("Use for Answers?") or _yes(get("Use for Answers?")),
                verification_status=status
                if (status := get("Verification Status"))
                in {"Verified", "Not verified", "Human Verification Required"}
                else "Not verified",
                accessibility="Public"
                if get("Accessibility") == "Public"
                else "Not human-verified",
                current_status=get("Current/Active Status") or "Unknown",
                authority_score=float(score) if re.fullmatch(r"\d+(\.\d+)?", score) else 0,
                found_by="Your team" if team_row else "Agent",
                # Agent rows were tool-returned in an earlier run; they are re-tested, not trusted.
                provenance="research_team" if team_row else "phase1_citation_sheet",
                notes=get("Notes"),
                team_notes=get("Your Notes"),
                checked_at=get("Last Checked"),
                content_path=text_path,
                snapshot_path=copy_path,
            )
        )
    assign_citation_ids(citations)  # rows the team added get the next free ID
    return parastatals, citations, issues


def write_sources_workbook(
    path: Path,
    city_name: str,
    parastatals: list[Parastatal],
    sources: list[Source],
    questions: dict[str, Question],
    applicability: dict[str, list[str]],
    issues: list[Issue],
) -> Path:
    """Write the Sources workbook, backing up any previous version to history/."""
    if path.exists():
        history = path.parent / "history"
        history.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, history / f"{path.stem}_{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    assign_citation_ids(sources)
    out_dir = path.parent
    wb = Workbook()
    guide = wb.active
    guide.title = "How to Use"
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 100
    for a, b in GUIDE:
        guide.append([a.format(city=city_name), b.format(city=city_name)])
    guide["A1"].font = Font(bold=True, size=14)
    for r in range(3, guide.max_row + 1):
        guide.cell(row=r, column=1).font = Font(bold=True)
        guide.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    ws = _new_sheet(
        wb,
        PARASTATALS_SHEET,
        PARASTATAL_COLUMNS,
        {
            "Parastatal Name": 40,
            "Why Included": 45,
            "Official Website": 35,
            "Website Check": 60,
            "Other Websites Checked": 45,
            "Governing Act": 30,
            "Notes": 50,
            "Your Notes": 30,
        },
    )
    rows = []
    for p in parastatals:
        check = p.website_check
        rows.append(
            (
                (p.id,),
                {
                    "Parastatal ID": p.id,
                    "Parastatal Name": p.name,
                    "Type": TYPE_LABELS[p.type],
                    "Include?": "Yes" if p.include else "No",
                    "Found By": p.found_by,
                    "Why Included": p.why_included,
                    "Official Website": Link(p.official_website)
                    if p.official_website
                    else "Not confirmed",
                    "Website Check": check.summary if check else "",
                    "Other Websites Checked": "\n".join(
                        f"{c.url}: {c.score}/100 ({c.band})" for c in p.other_websites
                    ),
                    "Governing Act": p.governing_act,
                    "Current Status": p.current_status,
                    **{
                        label: "Yes" if getattr(p, field) else "No"
                        for label, (field, _) in FLAGS.items()
                    },
                    "Notes": p.notes,
                    "Your Notes": p.team_notes,
                },
            )
        )
    mapping = _write_rows(ws, PARASTATAL_COLUMNS, rows, ["Parastatal ID"])
    for r, p in enumerate(parastatals, start=2):
        if p.website_check:
            colour = {"Official": "green", "Probably official": "amber"}.get(
                p.website_check.band, "red"
            )
            ws.cell(row=r, column=mapping["Website Check"]).fill = _FILLS[colour]
    _add_dropdown(ws, mapping["Include?"], ["Yes", "No"])
    _add_dropdown(ws, mapping["Type"], list(TYPES))
    for label in FLAGS:
        _add_dropdown(ws, mapping[label], ["Yes", "No"])

    ws = _new_sheet(
        wb,
        CITATIONS_SHEET,
        CITATION_COLUMNS,
        {
            "Source Title": 40,
            "Official URL": 45,
            "What This Source Is Useful For": 40,
            "Verification Status": 24,
            "Notes": 55,
            "Your Notes": 30,
        },
    )
    ordered = sorted(sources, key=lambda s: (s.parastatal_id, s.citation_id))
    mapping = _write_rows(
        ws,
        CITATION_COLUMNS,
        [
            (
                (s.citation_id,),
                {
                    "Citation ID": s.citation_id,
                    "Parastatal": s.parastatal_id,
                    "Source Title": s.title,
                    "Source Type": s.source_type,
                    "Official URL": Link(s.url),
                    "Saved Copy": _saved_copy(s.snapshot_path, out_dir),
                    "What This Source Is Useful For": s.useful_for,
                    "Question IDs": ", ".join(s.question_ids),
                    "Use for Answers?": "Yes" if s.use_for_answers else "No",
                    "Verification Status": s.verification_status,
                    "Accessibility": s.accessibility,
                    "Current/Active Status": s.current_status,
                    "Authority Score": s.authority_score,
                    "Found By": s.found_by,
                    "Notes": s.notes,
                    "Your Notes": s.team_notes,
                    "Last Checked": s.checked_at,
                },
            )
            for s in ordered
        ],
        ["Citation ID"],
    )
    _add_dropdown(ws, mapping["Use for Answers?"], ["Yes", "No"])

    ws = _new_sheet(
        wb, COVERAGE_SHEET, COVERAGE_COLUMNS, {"Question": 60, "Citations": 30, "Coverage": 40}
    )
    usable = [s for s in sources if s.use_for_answers and s.verification_status == "Verified"]
    rows = []
    for p in parastatals:
        if not p.include:
            continue
        for qid in applicability.get(p.id, []):
            q = questions[qid]
            if q.is_rollup:
                continue
            ids = [
                s.citation_id
                for s in usable
                if s.parastatal_id in (p.id, ALL_PARASTATALS) and q.id in s.question_ids
            ]
            rows.append(
                (
                    (p.id, q.raw_id),
                    {
                        "Parastatal": p.id,
                        "Question ID": q.raw_id,
                        "Question": q.text,
                        "Citations": ", ".join(ids),
                        "Coverage": "Has a source"
                        if ids
                        else "No source yet (add one if you know it)",
                    },
                )
            )
    mapping = _write_rows(ws, COVERAGE_COLUMNS, rows, ["Parastatal", "Question ID"])
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=mapping["Coverage"])
        cell.fill = _FILLS["green" if str(cell.value).startswith("Has") else "amber"]
    for cell in ws[1]:
        cell.fill = _HEADER_FILL

    _write_issue_sheets(wb, issues)
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:  # usually: the file is open in Excel on Windows
        path = path.with_name(f"{path.stem} (new {datetime.now():%H%M}).xlsx")
        wb.save(path)
    return path
